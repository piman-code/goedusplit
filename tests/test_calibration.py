import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from app.calibration import (
    borderline_indices, build_calibration, cut_lines, observed_target, observed_targets, suggest_presets,
    summary_lines, write_calibration_workbook,
)
from app.data_loader import ExamData, ItemInfo, StudentResponse

try:
    from app.main_window import MainWindow
except ModuleNotFoundError:
    MainWindow = None

CUTS = {"A": 80, "B": 60, "C": 50, "D": 20, "E": 10}


def _exam():
    """6 synthetic students. With 6 students, 2 are taken from each side of a cut.

    A cut 80: below [78, 70], above [82, 90]. B cut 60: below [52], above [70].
    C cut 50: below [48], above [52]. D and E: nobody within 10 points.
    """
    exam = ExamData(subject="합성")
    exam.cut_scores = dict(CUTS)
    exam.items = [
        ItemInfo(number=1, item_type="선택형", difficulty="쉬움", score=40.0, answer="1"),
        ItemInfo(number=2, item_type="선택형", difficulty="어려움", score=40.0, answer="2"),
        ItemInfo(number=1, item_type="서답형", difficulty="보통", score=20.0),
    ]
    data = [  # (final score, item1, item2, serdap score out of 20)
        (90, ".", ".", 20), (82, ".", "3", 16), (78, ".", "3", 12),
        (70, "2", ".", 10), (52, "2", "3", 4), (48, "2", "3", 0),
    ]
    for i, (score, a1, a2, serdap) in enumerate(data):
        exam.students.append(StudentResponse(
            sid=f"S{i}", class_no=f"합성반{i}", name=f"합성{i}", answers={1: a1, 2: a2},
            serdap_score=serdap, total=score, final_score=score,
        ))
    return exam, ["A", "A", "B", "B", "C", "미도달"]


def _design(kind, number, difficulty, points, rates, item):
    return {"type": kind, "number": number, "difficulty": difficulty, "target": "C", "points": points,
            "rates": dict(zip("ABCDE", rates)), "source_item": item}


def _standard_designs(exam):
    item1, item2, serdap = exam.items
    return [
        _design("선택형", 1, "쉬움", 40.0, (90, 80, 70, 60, 50), item1),
        _design("선택형", 2, "어려움", 40.0, (80, 60, 40, 20, 10), item2),
        _design("서답형", 1, "보통", 20.0, (70, 60, 50, 40, 30), serdap),
    ]


PRESETS = {  # the app's default table (보통 items)
    "A": (70, 45, 25, 10, 0), "B": (85, 70, 45, 25, 10), "C": (90, 80, 70, 45, 25),
    "D": (95, 90, 80, 70, 45), "E": (95, 95, 90, 80, 70),
}


def _rates_for(target, difficulty):
    delta = {"쉬움": 10, "어려움": -10}.get(difficulty, 0)
    return {lv: max(0, min(100, rate + delta)) for lv, rate in zip("ABCDE", PRESETS[target])}


def _border_row(difficulty, rates):
    return {"difficulty": difficulty, "border": dict(zip("ABCDE", rates))}


class ObservedTargetTests(unittest.TestCase):
    def test_target_is_the_closest_preset_row_for_the_difficulty(self):
        self.assertEqual(observed_target(_border_row("보통", (88, 78, 72, 50, 30)), _rates_for), "C")
        # 쉬움 items start 10 higher: the same pattern with +10 is still target C, not an easier level
        self.assertEqual(observed_target(_border_row("쉬움", (98, 88, 82, 60, 40)), _rates_for), "C")
        # a 쉬움 label with few correct at the lower cuts is a target A item
        self.assertEqual(observed_target(_border_row("쉬움", (75, 50, 30, 20, 15)), _rates_for), "A")

    def test_needs_three_boundaries_with_students(self):
        self.assertIsNone(observed_target(_border_row("보통", (90, 80, None, None, None)), _rates_for))
        self.assertEqual(observed_target(_border_row("보통", (95, 90, 80, None, None)), _rates_for), "D")

    def test_constructed_group_is_matched_with_its_items_difficulties(self):
        hard_b = _rates_for("B", "어려움")
        group = {"difficulty": "", "border": dict(hard_b), "members": [(10.0, "어려움"), (10.0, "어려움")]}
        self.assertEqual(observed_target(group, _rates_for), "B")
        # compared as a 보통 item, the same results would look like a target A group
        self.assertEqual(observed_target(dict(group, members=None), _rates_for), "A")
        exam, levels = _exam()
        report = build_calibration(_standard_designs(exam), exam, levels, rates_for=_rates_for)
        self.assertEqual(report["rows"][-1]["members"], [(20.0, "보통")])

    @unittest.skipIf(MainWindow is None, "main window unavailable")
    def test_with_the_apps_own_rules_each_default_row_is_its_own_target(self):
        window = MainWindow.__new__(MainWindow)  # default table, with the 2/3 and A≥…≥E rules applied
        for difficulty in ("쉬움", "보통", "어려움"):
            for target in "ABCDE":
                row = {"difficulty": difficulty, "border": dict(window._target_level_rates(target, difficulty))}
                self.assertEqual(observed_target(row, window._target_level_rates), target, (difficulty, target))

    def test_constructed_items_share_their_group_target(self):
        report = {"rows": [
            {"type": "선택형", "number": "2", "observed_target": "B"},
            {"type": "선택형", "number": "3", "observed_target": None},
            {"type": "서답형 묶음", "number": "1, 2", "observed_target": "A"},
        ]}
        self.assertEqual(observed_targets(report), {("선택형", 2): "B", ("서답형", 1): "A", ("서답형", 2): "A"})

    def test_report_names_items_whose_target_differs_from_the_data(self):
        exam, levels = _exam()
        report = build_calibration(_standard_designs(exam), exam, levels, rates_for=_rates_for)
        rows = {(r["type"], r["number"]): r for r in report["rows"]}
        # 쉬움 1번: A 75, B 0, C 0 / 어려움 2번: A 50, B 50, C 0 → both only A students pass: target A
        self.assertEqual([rows[key]["observed_target"] for key in (("선택형", "1"), ("선택형", "2"))], ["A", "A"])
        differ = [r for r in report["rows"] if r["type"] == "선택형" and r["observed_target"] != r["target"]]
        line = next(line for line in summary_lines(report) if line.startswith("목표수준이 실측과 다른"))
        self.assertIn(f"선택형 {len(differ)}문항", line)
        self.assertIsNone(build_calibration(_standard_designs(exam), exam, levels)["rows"][0]["observed_target"])


class CalibrationTests(unittest.TestCase):
    def test_borderline_takes_both_sides_of_each_cut(self):
        exam, _ = _exam()
        border = borderline_indices([st.final_score for st in exam.students], CUTS)
        self.assertEqual(border["A"], [[2, 3], [1, 0]])
        self.assertEqual(border["B"], [[4], [3]])
        self.assertEqual(border["C"], [[5], [4]])
        self.assertNotIn("D", border)
        # 한쪽에만 학생이 있으면 그쪽만, 동점은 함께 포함
        self.assertEqual(borderline_indices([55, 52, 52, 30], {"A": 60}), {"A": [[0, 1, 2]]})

    def test_rates_average_both_sides_and_relative_misses(self):
        exam, levels = _exam()
        report = build_calibration(_standard_designs(exam), exam, levels)
        rows = {(r["type"], r["number"]): r for r in report["rows"]}
        one, two, group = rows[("선택형", "1")], rows[("선택형", "2")], rows[("서답형 묶음", "1")]
        self.assertEqual([one["border"][lv] for lv in "ABC"], [75.0, 0.0, 0.0])      # A: 아래 50, 위 100
        self.assertEqual([two["border"][lv] for lv in "ABC"], [50.0, 50.0, 0.0])
        self.assertAlmostEqual(group["border"]["A"], 72.5)                             # 아래 55, 위 90
        self.assertAlmostEqual(group["border"]["B"], 35.0)
        self.assertIsNone(one["border"]["D"])
        self.assertEqual(one["all"]["A"], 100.0)
        self.assertEqual((one["diff"]["A"], two["diff"]["A"], group["diff"]["A"]), (-15.0, -30.0, 2.5))
        self.assertAlmostEqual(report["offset"]["A"], (40 * -15 + 40 * -30 + 20 * 2.5) / 100)
        self.assertAlmostEqual(report["offset"]["B"], -41.0)
        self.assertIsNone(report["offset"]["D"])
        # 상대차 = 차이 - 그 경계의 배점 가중 평균 차이
        self.assertEqual([one["relative"]["A"], two["relative"]["A"], group["relative"]["A"]], [2.5, -12.5, 20.0])
        self.assertEqual([one["relative"]["B"], two["relative"]["B"], group["relative"]["B"]], [-39.0, 31.0, 16.0])
        self.assertIsNone(one["relative"]["D"])
        self.assertEqual(report["large_gaps"], 5)  # A: 서답, B: 1·2번·서답, C: 1번
        easy, hard = report["relative_by_difficulty"]["쉬움"], report["relative_by_difficulty"]["어려움"]
        self.assertEqual((easy["items"], easy["levels"]["A"], easy["levels"]["C"]), (1, 2.5, -18.0))
        self.assertEqual(hard["levels"]["B"], 31.0)
        self.assertNotIn("", report["relative_by_difficulty"])  # 서답형 묶음은 난이도 없음
        cuts = {c["level"]: c for c in report["cuts"]}
        self.assertAlmostEqual(cuts["A"]["predicted_scaled"], 82.0)
        self.assertAlmostEqual(cuts["A"]["actual_scaled"], 64.5)
        # 배점 가중 평균 차이 = 분할점수 차이 (모든 문항에 실측이 있을 때)
        self.assertAlmostEqual(cuts["A"]["actual_scaled"] - cuts["A"]["predicted_scaled"], report["offset"]["A"])
        self.assertEqual(cuts["A"]["applied"], 80.0)
        self.assertIsNone(cuts["D"]["actual_scaled"])
        self.assertEqual(report["counts"], {"A": 2, "B": 2, "C": 1, "D": 0, "E": 0})
        self.assertEqual(report["border_counts"], {"A": 4, "B": 2, "C": 2, "D": 0, "E": 0})

    def test_serdap_group_is_point_weighted_and_needs_every_serdap_item(self):
        exam, levels = _exam()
        exam.items.append(ItemInfo(number=2, item_type="서답형", difficulty="보통", score=30.0))
        s1, s2 = exam.items[2], exam.items[3]
        both = [_design("서답형", 1, "보통", 10.0, (80, 70, 60, 50, 40), s1),
                _design("서답형", 2, "보통", 30.0, (40, 30, 20, 10, 0), s2)]
        report = build_calibration(both, exam, levels)
        self.assertEqual(report["rows"][0]["number"], "1, 2")
        self.assertAlmostEqual(report["rows"][0]["predicted"]["A"], (10 * 80 + 30 * 40) / 40)
        self.assertEqual(report["notes"], [])

        partial = build_calibration(both[:1] + [_design("서답형", 9, "보통", 5.0, (1, 1, 1, 1, 1), None)], exam, levels)
        self.assertEqual(partial["rows"], [])
        self.assertEqual(partial["unmatched"], ["서답형 9번"])
        self.assertIn("모든 서답형 문항", partial["notes"][0])

    def test_unmatched_undesigned_and_perform_notes(self):
        exam, levels = _exam()
        exam.use_perform, exam.weight_perform = True, 30.0
        designs = [_design("선택형", 1, "쉬움", 40.0, (90, 80, 70, 60, 50), exam.items[0]),
                   _design("선택형", 9, "보통", 10.0, (90, 80, 70, 60, 50), None)]
        report = build_calibration(designs, exam, levels)
        self.assertEqual(report["unmatched"], ["선택형 9번"])
        self.assertEqual(report["not_designed"], ["선택형 2번", "서답형 1번"])
        self.assertTrue(any("수행평가" in note for note in report["notes"]))

    def test_summary_reads_relative_misses_then_cautions(self):
        exam, levels = _exam()
        lines = summary_lines(build_calibration(_standard_designs(exam), exam, levels))
        self.assertIn("쉬움 문항(1개): 다른 문항보다 경계 학생 정답률을 평균 18.2%p 높게", lines[0])
        self.assertIn("어려움 문항(1개): 다른 문항보다 경계 학생 정답률을 평균 10.2%p 낮게", lines[1])
        self.assertIn("15%p 이상 빗나간 칸이 5개", lines[2])
        self.assertTrue(any(line.startswith("A/B: 경계 학생이 적음") for line in lines))
        self.assertTrue(any(line.startswith("D/E: 분할점수 가까이에 학생이 없어") for line in lines))
        report = build_calibration(_standard_designs(exam), exam, levels)
        report["border_one_sided"] = ["A"]
        self.assertTrue(any("한쪽에만" in line for line in summary_lines(report)))

    def test_cut_lines_compare_estimated_and_applied_cuts(self):
        exam, levels = _exam()
        lines = cut_lines(build_calibration(_standard_designs(exam), exam, levels))
        self.assertEqual(lines[0], "A/B: 검토안으로 계산한 분할점수 82.0점, 이번에 적용한 분할점수 80.0점.")
        self.assertEqual(len(lines), 5)
        partial = cut_lines(build_calibration(_standard_designs(exam)[:1], exam, levels))
        self.assertEqual(len(partial), 1)
        self.assertIn("일부 문항만", partial[0])
        exam.use_perform, exam.weight_perform = True, 30.0
        self.assertIn("수행평가", cut_lines(build_calibration(_standard_designs(exam), exam, levels))[0])

    def test_single_difficulty_group_gives_no_always_true_sentence(self):
        exam, levels = _exam()
        designs = _standard_designs(exam)
        for design in designs:
            design["difficulty"] = "보통"
        lines = summary_lines(build_calibration(designs, exam, levels))
        self.assertIn("난이도가 한 가지뿐", lines[0])
        self.assertFalse(any("비슷하게 예측" in line for line in lines))

    def test_preset_suggestion_removes_difficulty_adjustment(self):
        exam, levels = _exam()
        current = {t: {lv: 50 for lv in "ABCDE"} for t in "ABCDE"}
        designs = _standard_designs(exam)[:2]
        for design in designs:
            design["difficulty"] = "어려움"   # 두 문항 모두 어려움: 쓸 때 -10 되므로 제안은 +10
        suggestion = suggest_presets(build_calibration(designs, exam, levels), current)
        self.assertEqual(suggestion["C"]["suggested"]["A"], 73)  # (75+50)/2 + 10 = 72.5 → 73

    def test_preset_suggestion_uses_target_groups_and_keeps_missing_cells(self):
        exam, levels = _exam()
        report = build_calibration(_standard_designs(exam), exam, levels)
        current = {t: {lv: 10 * (i + 1) for i, lv in enumerate("ABCDE")} for t in "ABCDE"}
        suggestion = suggest_presets(report, current)
        # 목표 C: 1번(쉬움, 기준표에 +10 붙음) 실측 75/0/0 → 65/-10/-10, 2번(어려움, -10) 50/50/0 → 60/60/10
        # A (65+60)/2=62.5→63, B (-10+60)/2=25, C (-10+10)/2=0, D·E는 경계 학생이 없어 현재값 유지
        self.assertEqual(suggestion["C"]["suggested"], {"A": 63, "B": 25, "C": 0, "D": 40, "E": 50})
        self.assertEqual(suggestion["C"]["items"], 2)
        self.assertIsNone(suggestion["B"]["suggested"])
        self.assertIsNone(suggest_presets(report, current, min_items=3)["C"]["suggested"])

    def test_workbook_has_no_student_identity(self):
        import openpyxl
        exam, levels = _exam()
        with TemporaryDirectory() as directory:
            path = Path(directory) / "보정.xlsx"
            write_calibration_workbook(path, build_calibration(_standard_designs(exam), exam, levels))
            book = openpyxl.load_workbook(path)
            try:
                text = " ".join(str(c.value) for sheet in book for row in sheet.iter_rows() for c in row if c.value is not None)
                self.assertEqual(book.sheetnames, ["요약", "문항별"])
            finally:
                book.close()
        for student in exam.students:
            self.assertNotIn(student.name, text)
            self.assertNotIn(student.class_no, text)
            self.assertNotIn(student.sid, text)


def _judgment(rate):
    return {"correct": [], "targetRate": rate, "overrideRate": None}


@unittest.skipIf(MainWindow is None, "main window unavailable")
class CalibrationWindowTests(unittest.TestCase):
    def test_calculator_project_is_compared_with_judge_average(self):
        from types import SimpleNamespace
        from unittest.mock import patch
        window = MainWindow.__new__(MainWindow)
        window.exam, levels = _exam()
        window.overall = SimpleNamespace(levels_arr=levels)
        judges = [{"id": "j1", "name": "검토안 1"}, {"id": "j2", "name": "검토안 2"}]
        project = {"judges": judges, "items": [{
            "id": "i1", "number": 1, "type": "선택형", "difficulty": "쉬움", "targetLevel": "C", "points": 40,
            "judgmentsByJudge": {
                "j1": {lv: _judgment(r) for lv, r in zip("ABCDE", (90, 80, 70, 60, 50))},
                "j2": {lv: _judgment(r) for lv, r in zip("ABCDE", (70, 60, 50, 40, 30))},
            },
        }]}
        with patch.object(MainWindow, "_show_calibration_dialog") as show:
            window._on_calibration_project(project)
        report = show.call_args.args[0]
        self.assertEqual(report["rows"][0]["predicted"]["A"], 80.0)  # 두 검토안 평균
        self.assertEqual(report["rows"][0]["border"]["A"], 75.0)
        self.assertEqual(report["not_designed"], ["선택형 2번", "서답형 1번"])

    def _import_window(self):
        from types import SimpleNamespace
        window = MainWindow.__new__(MainWindow)
        window.exam, levels = _exam()
        window.overall = SimpleNamespace(levels_arr=levels)
        return window

    @staticmethod
    def _paper_rows(points=(40.0, 40.0, 20.0)):
        return [{"type": kind, "number": number, "points": score, "difficulty": difficulty, "flags": [], "choices": 5, "preview": ""}
                for (kind, number, difficulty), score in zip((("선택형", 1, "쉬움"), ("선택형", 2, "어려움"), ("서답형", 1, "보통")), points)]

    def test_import_takes_targets_from_the_analysis_of_the_same_paper(self):
        window = self._import_window()
        targets, analysis = window._observed_import_targets(self._paper_rows())
        self.assertEqual(targets, {("선택형", 1): "A", ("선택형", 2): "A", ("서답형", 1): "A"})
        self.assertEqual(analysis, "합성 · 6명")
        # 배점 the paper did not show were filled from this very analysis: they cannot count as agreement
        rows = self._paper_rows()
        paper = {(r["type"], r["number"]): (40.0 if r["number"] == 1 and r["type"] == "선택형" else None) for r in rows}
        self.assertEqual(set(window._observed_import_targets(rows, paper)[0]), {("선택형", 1), ("선택형", 2), ("서답형", 1)})
        self.assertEqual(window._observed_import_targets(rows, {key: None for key in paper}), ({}, ""))
        # another paper (배점 differ) or another item count: no targets from this analysis
        self.assertEqual(window._observed_import_targets(self._paper_rows((30.0, 50.0, 20.0))), ({}, ""))
        self.assertEqual(window._observed_import_targets(self._paper_rows()[:2]), ({}, ""))

    def test_paper_must_name_subject_year_and_semester_of_the_analysis(self):
        window = self._import_window()
        window.exam.subject, window.exam.semester = "공통수학1(4)", "2026학년도 1학기"
        self.assertTrue(window._paper_names_loaded_exam("2026학년도 1학기 1차 지필평가 공통수학1 1학년"))
        self.assertTrue(window._paper_names_loaded_exam("2026 학년도 1 학기 … 공통 수학1"))
        for text in ("2026학년도 1학기 공통수학1Ⅱ", "2026학년도 2학기 공통수학1", "2025학년도 1학기 공통수학1",
                     "2026학년도 11학기 공통수학1", "2026학년도 1학기 수학"):
            self.assertFalse(window._paper_names_loaded_exam(text), text)
        window.exam.subject = "수학(4)"
        self.assertFalse(window._paper_names_loaded_exam("2026학년도 1학기 공통수학1"))  # a longer name is another subject
        window.exam.semester = ""
        self.assertFalse(window._paper_names_loaded_exam("2026학년도 1학기 수학"))

    def test_import_uses_analysis_targets_by_default_only_when_the_paper_names_the_exam(self):
        import os
        from unittest.mock import patch
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication
        self.app = QApplication.instance() or QApplication([])  # the import shows a wait cursor
        window = self._import_window()
        window.exam.subject, window.exam.semester = "합성(4)", "2026학년도 1학기"  # NEIS adds the credits
        parsed = {"items": self._paper_rows(), "total_points": 100.0, "warnings": []}
        for text, expected in (("2026학년도 1학기 합성 시험", True), ("2026학년도 1학기 다른 과목 시험", False)):
            with self.subTest(text=text), \
                 patch("app.main_window.QFileDialog.getOpenFileName", return_value=("paper.hwp", "")), \
                 patch("app.main_window.extract_exam_text", return_value=(text, "합성")), \
                 patch("app.main_window.parse_exam_structure", return_value=dict(parsed, items=[dict(r) for r in parsed["items"]])), \
                 patch.object(MainWindow, "_show_exam_structure_dialog") as show:
                window.import_exam_structure()
            kwargs = show.call_args.kwargs
            self.assertEqual(kwargs["use_observed"], expected)
            self.assertEqual(kwargs["analysis"], "합성(4) · 2026학년도 1학기 · 6명")
            rows = show.call_args.args[0]["items"]
            chosen = [row["observed_target"] if expected else row["rule_target"] for row in rows]
            self.assertEqual([row["target"] for row in rows], chosen)

    def test_preview_checkbox_switches_targets_but_keeps_the_teachers_own_choice(self):
        import os
        from unittest.mock import patch
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication, QCheckBox, QComboBox, QDialog, QWidget
        self.app = QApplication.instance() or QApplication([])
        owner = QWidget()
        owner._px = lambda value: int(value)
        owner._import_target_level = MainWindow._import_target_level
        rows = [dict(row, observed_target=target, standard="") for row, target in zip(self._paper_rows(), "ADA")]
        seen = {}

        def inspect(dialog):
            check = dialog.findChildren(QCheckBox)[0]
            targets = [box for box in dialog.findChildren(QComboBox) if box.count() == 5]
            difficulties = [box for box in dialog.findChildren(QComboBox) if box.count() == 3]
            seen["on"] = [box.currentText() for box in targets]
            check.setChecked(False)
            seen["off"] = [box.currentText() for box in targets]
            targets[1].setCurrentIndex(4)
            targets[1].activated.emit(4)            # the teacher picks E for 선택형 2번
            difficulties[1].setCurrentText("쉬움")  # and changes its difficulty
            check.setChecked(True)
            seen["mixed"] = [box.currentText() for box in targets]
            return 0

        with patch.object(QDialog, "exec", inspect):
            MainWindow._show_exam_structure_dialog(owner, {"items": rows, "total_points": 100.0, "warnings": []},
                                                   "합성", "paper.hwp", analysis="합성 · 6명", use_observed=True)
        self.assertEqual(seen["on"], ["A", "D", "A"])
        self.assertEqual(seen["off"], ["E", "B", "C"])       # difficulty rule
        self.assertEqual(seen["mixed"], ["A", "E", "A"])     # 2번 keeps the teacher's E

    def test_applying_suggestion_changes_only_chosen_rows_and_keeps_rules(self):
        import json
        from unittest.mock import Mock
        from app.main_window import DEFAULT_TARGET_RATE_PRESETS
        window = MainWindow.__new__(MainWindow)
        saved = {}
        window.settings = Mock(setValue=lambda key, value: saved.__setitem__(key, value))
        window._send_spliter_teacher_presets = Mock()
        window.statusBar = Mock()
        current = {t: dict(DEFAULT_TARGET_RATE_PRESETS[t]) for t in "ABCDE"}
        suggestion = {t: {"items": 0, "current": current[t], "suggested": None} for t in "ABCDE"}
        suggestion["C"]["suggested"] = {"A": 63, "B": 25, "C": 40, "D": 20, "E": 5}   # C < 2/3 → 규칙으로 보정
        suggestion["B"]["suggested"] = {"A": 90, "B": 80, "C": 50, "D": 30, "E": 10}  # 고르지 않은 행
        window._apply_preset_suggestion(suggestion, ["C"])
        stored = json.loads(saved["spliter/target_rate_presets"])
        self.assertEqual(stored["B"], current["B"])
        self.assertNotEqual(stored["C"], current["C"])
        self.assertGreaterEqual(stored["C"]["C"], 66)                 # 목표수준 학생 2/3 이상
        rates = [stored["C"][lv] for lv in "ABCDE"]
        self.assertEqual(rates, sorted(rates, reverse=True))         # A ≥ B ≥ … ≥ E
        window._send_spliter_teacher_presets.assert_called_once_with(apply_current=False)

    def _sample_project(self):
        target = lambda n: "E" if n <= 3 else "D" if n <= 6 else "C" if n <= 11 else "B" if n <= 15 else "A"
        difficulty = lambda n: {"A": "어려움", "B": "어려움", "C": "보통"}.get(target(n), "쉬움")
        points = (4, 4, 5, 5, 5, 5, 5, 5, 5, 5, 6, 6, 6, 6, 6, 6, 8, 8)
        return {"items": [{"number": n, "type": "선택형", "title": f"{n}번", "targetLevel": target(n), "difficulty": difficulty(n),
                           "standard": "", "note": "", "evidence": [], "points": points[n - 1]} for n in range(1, 19)]}

    def test_calculator_example_items_are_recognised(self):
        sample = self._sample_project()
        self.assertTrue(MainWindow._is_sample_calculator_project(sample))
        for change in ({"title": "1번 다항식"}, {"standard": "[10수학01-01]"}, {"note": "메모"}, {"targetLevel": "C"},
                       {"points": 3.5}, {"difficulty": "보통"}):
            edited = self._sample_project()
            edited["items"][0].update(change)
            self.assertFalse(MainWindow._is_sample_calculator_project(edited), change)
        shorter = self._sample_project()
        shorter["items"].pop()
        self.assertFalse(MainWindow._is_sample_calculator_project(shorter))
        self.assertFalse(MainWindow._is_sample_calculator_project(None))

    def test_example_items_need_confirmation_before_comparing(self):
        from unittest.mock import patch
        window = MainWindow.__new__(MainWindow)
        window.exam, _ = _exam()
        with patch("app.main_window.QMessageBox") as messages, \
             patch.object(MainWindow, "_show_calibration_dialog") as show:
            messages.question.return_value = messages.No
            window._on_calibration_project(self._sample_project())
        messages.question.assert_called_once()
        show.assert_not_called()

    def test_shown_preset_row_is_what_gets_saved(self):
        import json, random
        from unittest.mock import Mock
        window = MainWindow.__new__(MainWindow)
        saved = {}
        window.settings = Mock(setValue=lambda key, value: saved.__setitem__(key, value))
        rng = random.Random(3)
        for _ in range(3000):
            target = rng.choice("ABCDE")
            row = window._settled_preset_row(target, {lv: rng.randint(0, 100) for lv in "ABCDE"})
            window._save_target_rate_presets({target: row})
            self.assertEqual(json.loads(saved["spliter/target_rate_presets"])[target], row)

    def test_empty_calculator_and_missing_analysis_explain_next_step(self):
        from unittest.mock import patch
        window = MainWindow.__new__(MainWindow)
        window.exam, window.overall = None, None
        with patch("app.main_window.QMessageBox") as messages:
            window.show_calibration_report()
            window.exam, _ = _exam()
            window._on_calibration_project({"items": []})
        texts = [call.args[2] for call in messages.information.call_args_list]
        self.assertIn("분석을 실행", texts[0])
        self.assertIn("작업 불러오기", texts[1])


if __name__ == "__main__":
    unittest.main()
