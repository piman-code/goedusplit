from tests.runtime_isolation import ensure_isolated
ensure_isolated()
import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import openpyxl
from PySide6.QtWidgets import QApplication
from app.folder_batch import discover_round_groups
from app.main_window import MainWindow
from app.synthetic_qa import analysis_fixtures
from app.work_session_ui import capture_ui_state, apply_python_state
from app.work_sessions import build_session, prepare_session, WorkSessionStore
from app.portfolio_details import load_work_session_rows, resolve_details


def four_round_fixtures(root):
    original=root/'original';original.mkdir()
    folder=root/'four';folder.mkdir()
    item,responses,performs=analysis_fixtures(original)
    for semester in (1,2):
        for number in (1,2):
            book=openpyxl.load_workbook(item)
            book.save(folder/f'2026학년도 {semester}학기 {number}차 문항정보표.xlsx')
            for klass,path in enumerate(responses[number],1):
                book=openpyxl.load_workbook(path)
                book.active.cell(3,1,f'2026학년도 {semester}학기 {number}차 1학년 수학:합성수학')
                # Each round has a real distinct score, preserved through encoding.
                book.active.cell(7,6,semester*20+number*5)
                book.active.cell(7,9,semester*20+number*5)
                book.save(folder/f'{semester}학기 {number}차 정오표(1-{klass}).xlsx')
        for klass,path in enumerate(performs,1):
            book=openpyxl.load_workbook(path)
            book.save(folder/f'{semester}학기 수행평가(1-{klass}).xlsx')
    return folder

class FolderRoundBatchTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    def setUp(self):
        self.root=Path(tempfile.mkdtemp(prefix='folder-batch-'))
        self.folder=four_round_fixtures(self.root)
        self.web=patch('app.main_window.QWebEngineView',None);self.web.start();self.addCleanup(self.web.stop)
        self.rootpatch=patch.object(MainWindow,'_ai_material_root_dir',lambda w:self.root/'state');self.rootpatch.start();self.addCleanup(self.rootpatch.stop)
        self.w=MainWindow();self.addCleanup(self.w.close)
        self.w.spin_pencil_ratio.setValue(60);self.w.spin_perform_ratio.setValue(40)
    def test_four_rounds_real_values_and_full_save_fresh_restore(self):
        groups=discover_round_groups(self.folder,'1학기 1차')
        self.assertEqual([(1,1),(1,2),(2,1),(2,2)],[g['key'][1:3] for g in groups])
        self.assertTrue(all(not g['problems'] and len(g['perform'])==2 for g in groups))
        original={p:p.read_bytes() for p in self.folder.glob('*.xlsx')}
        with patch('app.main_window.QMessageBox.warning') as warning:
            self.assertTrue(self.w._analyze_folder_groups(groups),warning.call_args)
        self.assertEqual(4,len(self.w._portfolio_analysis_sessions))
        self.assertEqual(4,len(self.w._round_store().load_all()))
        store=WorkSessionStore(self.root/'full')
        saved=store.save(build_session(self.w.exam,self.w.perform_data,capture_ui_state(self.w),None))
        fresh=MainWindow();self.addCleanup(fresh.close)
        apply_python_state(fresh,prepare_session(store.read(saved)))
        rows=[r for r in load_work_session_rows(fresh) if r['class_no']=='1/1']
        # The current fourth is represented by the current analysis, the prior three by cache.
        self.assertEqual(3,len(rows))
        self.assertEqual(3,len({r['score'] for r in rows}))
        for row in rows:
            self.assertTrue(all(tab['available'] for tab in resolve_details(fresh,row)['tabs'].values()))
        self.assertEqual(original,{p:p.read_bytes() for p in original})
    def test_all_preflight_failure_leaves_existing_inputs_and_domain(self):
        groups=discover_round_groups(self.folder,'1학기 1차')
        self.assertTrue(self.w._analyze_folder_groups(groups[:1]))
        before=capture_ui_state(self.w);exam=self.w.exam
        groups[-1]['iteminfo']=str(self.root/'does-not-exist.xlsx')
        with patch('app.main_window.QMessageBox.warning'):
            self.assertFalse(self.w._analyze_folder_groups(groups))
        self.assertIs(exam,self.w.exam)
        self.assertEqual(before,capture_ui_state(self.w))
    def test_missing_round_item_requires_explicit_common_selection(self):
        folder=self.root/'legacy';folder.mkdir()
        source=next(self.folder.glob('*1학기 1차 문항정보표.xlsx'))
        (folder/'문항정보표(합성수학) 샘플.xlsx').write_bytes(source.read_bytes())
        source=next(self.folder.glob('1학기 1차 정오표*.xlsx'))
        book=openpyxl.load_workbook(source);book.active.cell(3,1,'2026학년도 1학년 수학:합성수학')
        book.save(folder/'1차 정기시험 교과목별 학생답 정오표(1-1).xlsx')
        groups=discover_round_groups(folder,'2학기 2차')
        self.assertEqual((2,1),groups[0]['key'][1:3])
        self.assertFalse(groups[0]['iteminfo'])
        self.assertEqual(1,len(groups[0]['item_candidates']))
        self.assertFalse(groups[0]['problems'])
        self.assertTrue(groups[0]['warnings'])
    def test_duplicate_item_is_not_chosen_by_mtime(self):
        source=next(self.folder.glob('*1학기 1차 문항정보표.xlsx'))
        (self.folder/'2026학년도 1학기 1차 문항정보표 사본.xlsx').write_bytes(source.read_bytes())
        group=discover_round_groups(self.folder,'1학기 1차')[0]
        self.assertFalse(group['iteminfo'])
        self.assertTrue(any('중복' in message for message in group['problems']))

    def test_header_conflict_and_other_subject_do_not_join(self):
        path=next(self.folder.glob('1학기 1차 정오표*.xlsx'))
        book=openpyxl.load_workbook(path)
        book.active.cell(3,1,'2026학년도 2학기 1차 1학년 수학:합성수학')
        book.save(path)
        groups=discover_round_groups(self.folder,'1학기 1차')
        self.assertTrue(any('정오표 안 시험 정보 충돌' in g['problems'] for g in groups))
        for group in groups:
            for perform in group['perform']:
                self.assertIn(f"{group['key'][1]}학기",Path(perform).name)

    def test_actual_style_names_read_header_and_do_not_auto_share_sample(self):
        folder=self.root/'actual-style';folder.mkdir()
        sources=self.root/'style-source';sources.mkdir()
        item,responses,performs=analysis_fixtures(sources)
        (folder/'문항정보표(합성수학) 샘플.xlsx').write_bytes(item.read_bytes())
        for number in (1,2):
            for klass,path in enumerate(responses[number],1):
                (folder/f'{number}차 정기시험 교과목별 학생답 정오표(1-{klass}).xlsx').write_bytes(path.read_bytes())
        for klass,path in enumerate(performs,1):
            (folder/f'수행평가 강의실별 일람표(1-{klass}).xlsx').write_bytes(path.read_bytes())
        groups=discover_round_groups(folder,'2학기 2차')
        self.assertEqual([(1,1),(1,2)],[g['key'][1:3] for g in groups])
        self.assertTrue(all(len(g['response'])==2 and not g['iteminfo'] for g in groups))
        self.assertTrue(all(len(g['item_candidates'])==1 and len(g['perform_candidates'])==2 for g in groups))
        # A teacher explicitly assigns the unlabelled shared items/performance to each round.
        for group in groups:
            group['iteminfo']=group['item_candidates'][0]
            group['perform']=group['perform_candidates']
            group['problems']=[]
        self.assertTrue(self.w._analyze_folder_groups(groups))
        self.assertTrue(self.w.exam.use_perform)
        self.assertEqual(2,len(self.w._portfolio_analysis_sessions))

    def test_batch_uses_same_explicit_weights_as_single_analysis(self):
        from app.folder_batch import prepare_round_groups
        from app.data_loader import load_exam, apply_perform, PATH_SEPARATOR
        from app.perform_loader import load_perform_many
        group=discover_round_groups(self.folder,'1학기 1차')[0]
        exam=load_exam(group['iteminfo'],PATH_SEPARATOR.join(group['response']))
        perform=load_perform_many(PATH_SEPARATOR.join(group['perform']))
        apply_perform(exam,perform,70,30)
        actual=prepare_round_groups([group],exam.cut_scores,70,30)[0][1]
        self.assertEqual((70,30),(actual.weight_pencil,actual.weight_perform))
        self.assertEqual([s.final_score for s in exam.students],[s.final_score for s in actual.students])

    def test_explicit_conflicting_auxiliary_cannot_be_forced_into_group(self):
        from app.folder_batch import prepare_round_groups
        group=discover_round_groups(self.folder,'1학기 1차')[0]
        source=Path(group['iteminfo'])
        for label in ('2027학년도 1학기 1차 문항정보표','2026학년도 1학기 1차 2학년 문항정보표',
                      '2026학년도 1학기 1차 문항정보표(다른과목)'):
            with self.subTest(label=label):
                path=self.root/(label+'.xlsx');path.write_bytes(source.read_bytes())
                changed=copy.deepcopy(group);changed['iteminfo']=str(path)
                with self.assertRaises(ValueError):prepare_round_groups([changed],dict(A=90,B=80,C=70,D=60,E=40),60,40)

    def test_first_round_only_does_not_create_other_rounds(self):
        folder=self.root/'one';folder.mkdir()
        for source in self.folder.glob('*1학기 1차*.xlsx'):
            (folder/source.name).write_bytes(source.read_bytes())
        self.assertEqual([(1,1)],[g['key'][1:3] for g in discover_round_groups(folder,'2학기 2차')])

    def test_new_batch_replaces_stale_summary_selection_with_analysis(self):
        groups=discover_round_groups(self.folder,'1학기 1차')
        self.assertTrue(self.w._analyze_folder_groups(groups[:1]))
        self.w.save_current_subject_snapshot()
        self.w.refresh_portfolio_tab()
        key=next(iter(self.w._current_identity_by_hash()))
        self.w.combo_portfolio_student.setCurrentIndex(self.w.combo_portfolio_student.findData('hash:'+key))
        self.w.combo_portfolio_subject.setCurrentIndex(self.w.combo_portfolio_subject.findData(self.w.exam.subject))
        view=self.w.portfolio_view;view.sync()
        view.round_combo.setCurrentIndex(0)
        summary=next(i for i,row in enumerate(view._versions) if row.get('source_kind') not in ('current','analysis-session','work-session'))
        view.version_combo.setCurrentIndex(summary)
        self.assertFalse(resolve_details(self.w,view.selected_point)['tabs']['items']['available'])
        before=capture_ui_state(self.w)
        broken=copy.deepcopy(groups);broken[-1]['iteminfo']=str(self.root/'missing.xlsx')
        with patch('app.main_window.QMessageBox.warning'):
            self.assertFalse(self.w._analyze_folder_groups(broken))
        self.assertEqual(before,capture_ui_state(self.w))
        self.assertTrue(self.w._analyze_folder_groups(groups))
        for index in range(4):
            view.round_combo.setCurrentIndex(index)
            self.assertTrue(all(t['available'] for t in resolve_details(self.w,view.selected_point)['tabs'].values()),
                            (index,view.selected_point.get('source_kind')))

    def test_dialog_default_four_selection_accept_and_cancel_preserve_work(self):
        from PySide6.QtCore import QTimer
        from PySide6.QtWidgets import QDialog,QCheckBox
        groups=discover_round_groups(self.folder,'1학기 1차')
        before=capture_ui_state(self.w)
        seen=[]
        def close_dialog(accept):
            dialog=QApplication.activeModalWidget()
            if not isinstance(dialog,QDialog):
                QTimer.singleShot(10,lambda:close_dialog(accept));return
            checks=[c for c in dialog.findChildren(QCheckBox) if c.text() in {g['label'] for g in groups}]
            seen.append((len(checks),all(c.isChecked() for c in checks)))
            dialog.accept() if accept else dialog.reject()
        QTimer.singleShot(0,lambda:close_dialog(False))
        self.assertEqual([],self.w._choose_folder_rounds(groups))
        self.assertEqual(before,capture_ui_state(self.w))
        QTimer.singleShot(0,lambda:close_dialog(True))
        self.assertEqual(4,len(self.w._choose_folder_rounds(groups)))
        self.assertEqual([(4,True),(4,True)],seen)
        self.assertEqual(before,capture_ui_state(self.w))

    def test_optional_numeric_subject_annotation_preserves_identity_and_four_rounds(self):
        from app.folder_batch import _metadata
        from app.data_loader import load_exam,PATH_SEPARATOR
        for path in self.folder.glob('*정오표*.xlsx'):
            book=openpyxl.load_workbook(path)
            book.active.cell(3,1,book.active.cell(3,1).value.replace('합성수학','합성수학(4)'))
            book.save(path)
        groups=discover_round_groups(self.folder,'1학기 1차')
        self.assertEqual(4,len(groups))
        self.assertTrue(all(g['key'][4]=='합성수학(4)' and not g['problems'] for g in groups))
        self.assertTrue(all(len(g['perform'])==2 for g in groups))
        expected=load_exam(groups[-1]['iteminfo'],PATH_SEPARATOR.join(groups[-1]['response']))
        self.assertEqual('합성수학(4)',expected.subject)
        self.assertTrue(self.w._analyze_folder_groups(groups))
        self.assertEqual(expected.subject,self.w.exam.subject)
        identity=self.w._current_identity_by_hash()
        snapshot=self.w._current_subject_snapshot()
        self.assertEqual(expected.subject,snapshot['subject'])
        store=WorkSessionStore(self.root/'qualified-work')
        path=store.save(build_session(self.w.exam,self.w.perform_data,capture_ui_state(self.w),None))
        fresh=MainWindow();self.addCleanup(fresh.close)
        fresh.settings.setValue('privacy/portfolio_hash_key',self.w.settings.value('privacy/portfolio_hash_key'))
        apply_python_state(fresh,prepare_session(store.read(path)))
        self.assertEqual(identity,fresh._current_identity_by_hash())
        self.assertEqual('합성수학(4)',fresh.exam.subject)
        self.assertEqual(4,len(fresh._portfolio_analysis_sessions))
        # Base filename plus qualified workbook title remains one qualified subject.
        item=self.root/'문항정보표(합성수학).xlsx'
        book=openpyxl.load_workbook(groups[0]['iteminfo'])
        book.active.cell(1,1,'합성 검증 전용 (합성수학(4)) 과목');book.save(item)
        self.assertEqual({'합성수학(4)'},_metadata(item)['과목'])

    def test_subject_annotation_compatibility_does_not_merge_explicit_differences(self):
        from app.folder_batch import subjects_compatible,prepare_round_groups
        self.assertTrue(subjects_compatible('공통수학1(4)','공통수학1'))
        self.assertTrue(subjects_compatible('공통수학1','공통수학1(4)'))
        for left,right in [('공통수학1(4)','공통수학1(3)'),('공통수학1','공통수학2'),
                           ('공통수학1(심화)','공통수학1'),('공통수학1(4)','공통수학2(4)')]:
            self.assertFalse(subjects_compatible(left,right),(left,right))
        group=discover_round_groups(self.folder,'1학기 1차')[0]
        group['key']=(*group['key'][:4],'합성수학(4)')
        path=self.root/'1학기 1차 문항정보표(합성수학(3)).xlsx'
        book=openpyxl.load_workbook(group['iteminfo'])
        book.active.cell(1,1,'합성 검증 전용 (합성수학(3)) 과목');book.save(path)
        group['iteminfo']=str(path)
        with self.assertRaises(ValueError):prepare_round_groups([group],dict(A=90,B=80,C=70,D=60,E=40),60,40)

    def test_bare_auxiliary_for_two_explicit_numeric_subjects_requires_selection(self):
        from app.folder_batch import subjects_compatible
        import unicodedata
        self.assertTrue(subjects_compatible(unicodedata.normalize('NFD','공통수학1 （4）'),'공통수학1'))
        self.assertTrue(subjects_compatible('공통수학1（4）','공통수학1(4)'))
        for path in self.folder.glob('*정오표*.xlsx'):
            book=openpyxl.load_workbook(path)
            suffix='3' if '1학기' in path.name else '4'
            book.active.cell(3,1,book.active.cell(3,1).value.replace('합성수학',f'합성수학({suffix})'))
            book.save(path)
        groups=discover_round_groups(self.folder,'1학기 1차')
        self.assertEqual({'합성수학(3)','합성수학(4)'},{g['key'][4] for g in groups})
        self.assertTrue(all(not g['iteminfo'] and len(g['item_candidates'])==1 for g in groups))
        self.assertTrue(all(not g['perform'] and len(g['perform_candidates'])==2 for g in groups))
