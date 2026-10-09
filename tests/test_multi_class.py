from tests.runtime_isolation import ensure_isolated
ensure_isolated()

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

import openpyxl

from app.analysis import analyze_overall, build_score_matrix
from app.data_loader import PATH_SEPARATOR, apply_perform, load_exam, load_student_responses, split_paths
from app.perform_loader import load_perform_many

LEGEND = "※ . : 맞음 , 번호 : 틀림 , 알파벳 : 복수답안코드 (합성 범례)"


def _item_info(path):
    book = openpyxl.Workbook()
    sheet = book.active
    for row in [["합성 검증 전용 (합성수학) 과목"], ["선택형 문항"],
                ["문항번호", "내용영역", "성취기준", "난이도", None, None, "배점", "정답"],
                [None, None, None, "어려움", "보통", "쉬움"],
                [1, "합성 연산", "[합성-01] 합성 기준", None, None, "○", 50, 1],
                [2, "합성 연산", "[합성-02] 합성 기준", None, "○", None, 50, 3]]:
        sheet.append(row)
    book.save(path)


def _responses(path, klass, students, numeric=False, legend=True):
    """students: [(number, answer1, answer2, total)]. numeric=True stores choices as Excel numbers (3.0)."""
    def cell(value):
        return float(value) if numeric and value not in (".", "") else value
    book = openpyxl.Workbook()
    sheet = book.active
    for row in [["합성 검증 전용"], [], ["2026학년도 1학기 1학년 수학:합성수학"],
                ["반/번호", None, "성명", 1, 2, "선택형점수", "서답형점수", "기타점수", "과목총점"],
                [None, None, "정답", cell("1"), cell("3")], [None, None, "배점", 50, 50]]:
        sheet.append(row)
    for number, a1, a2, total in students:
        sheet.append([f"2026{klass}{number:03d}", f"{klass}/{number}", f"합성{klass}{number}",
                      cell(a1), cell(a2), total, 0, 0, total])
    if legend:
        sheet.append([LEGEND])
    book.save(path)


def _perform(path, klass, scores, area_max=40):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["합성 검증 전용"])
    sheet.append(["교과목 : 합성수학"])
    sheet.append(["반/번호", "합성 학번", "성명", f"합성 영역(만점 {area_max}.00,40.00%)", "합 계"])
    for number, score in scores:
        sheet.append([f"{klass}/{number}", f"2026{klass}{number:03d}", f"합성{klass}{number}", score, score])
    book.save(path)


class MultiClassTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        d = Path(self.tmp.name)
        self.item, self.c1, self.c2, self.c2_other = d / "items.xlsx", d / "정오표(1-1).xlsx", d / "정오표(1-2).xlsx", d / "다른.xlsx"
        _item_info(self.item)
        _responses(self.c1, 1, [(1, ".", ".", 100), (2, ".", "2", 50)])
        _responses(self.c2, 2, [(1, "4", ".", 50), (2, ".", ".", 100), (3, "4", "2", 0)], numeric=True)
        _responses(self.c2_other, 2, [(1, ".", ".", 100)])
        wb = openpyxl.load_workbook(self.c2_other)       # a different answer key
        wb.active["D5"] = "2"
        wb.save(self.c2_other)
        self.p1, self.p2 = d / "수행(1-1).xlsx", d / "수행(1-2).xlsx"
        _perform(self.p1, 1, [(1, 40), (2, 20)])
        _perform(self.p2, 2, [(1, 10), (2, 30), (3, 0)])

    def test_choices_stored_as_numbers_read_like_text(self):
        students, answers, _ = load_student_responses(self.c2)
        self.assertEqual({1: "1", 2: "3"}, answers)                  # not '1.0' / '3.0'
        self.assertEqual("4", students[0].answers[1])
        exam = load_exam(self.item, self.c2)
        _score, choice, _items = build_score_matrix(exam)
        self.assertEqual([4, 3], choice[0].tolist())                 # wrong pick 4, right pick 3 (was 0 before)

    def test_footnote_line_under_the_table_is_not_a_student(self):
        students, _answers, _ = load_student_responses(self.c1)
        self.assertEqual(2, len(students))
        self.assertTrue(all(s.sid.isdigit() for s in students))

    def test_several_class_files_become_one_exam(self):
        exam = load_exam(self.item, [self.c1, self.c2])
        self.assertEqual(5, len(exam.students))
        self.assertEqual({"1", "2"}, {s.grade_class for s in exam.students})
        self.assertEqual(PATH_SEPARATOR.join([str(self.c1), str(self.c2)]), exam.source_files["responses"])
        text_form = load_exam(self.item, PATH_SEPARATOR.join([str(self.c1), str(self.c2)]))
        self.assertEqual(5, len(text_form.students))
        perform = load_perform_many([self.p1, self.p2])
        apply_perform(exam, perform, 60, 40)
        self.assertEqual(5, sum(1 for s in exam.students if s.sid in perform.records))
        self.assertEqual({"1", "2"}, set(analyze_overall(exam).by_class))
        single = load_exam(self.item, self.c1)
        self.assertEqual(2, len(single.students))                    # one file behaves as before

    def test_mixing_up_files_is_refused_with_a_reason(self):
        with self.assertRaisesRegex(ValueError, "모두 있습니다"):
            load_exam(self.item, [self.c1, self.c1])
        with self.assertRaisesRegex(ValueError, "정답이 .* 다릅니다"):
            load_exam(self.item, [self.c1, self.c2_other])
        with self.assertRaisesRegex(ValueError, "모두 있습니다"):
            load_perform_many([self.p1, self.p1])
        other = Path(self.tmp.name) / "수행-다른.xlsx"
        _perform(other, 2, [(1, 10)], area_max=50)
        with self.assertRaisesRegex(ValueError, "영역·만점·반영비율"):
            load_perform_many([self.p1, other])

    def test_split_paths(self):
        self.assertEqual(["a.xlsx", "b.xlsx"], split_paths("a.xlsx|b.xlsx"))
        self.assertEqual(["a.xlsx"], split_paths(" a.xlsx "))
        self.assertEqual(["a", "b"], split_paths(["a", "b"]))


class MultiClassWindowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        from PySide6.QtCore import QSettings
        from app.main_window import MainWindow
        self.resources = ExitStack()
        self.addCleanup(self.resources.close)
        tmp = Path(self.resources.enter_context(tempfile.TemporaryDirectory()))
        settings = QSettings(str(tmp / "settings.ini"), QSettings.IniFormat)
        self.resources.enter_context(patch("app.main_window.QSettings", return_value=settings))
        self.resources.enter_context(patch("app.main_window.QWebEngineView", None))
        self.resources.enter_context(patch.object(MainWindow, "_ai_material_root_dir", lambda _self: tmp / "appdata"))
        self.window = MainWindow()
        self.addCleanup(self.window.close)

    def test_file_dialog_can_take_several_class_files(self):
        from PySide6.QtWidgets import QFileDialog
        with patch.object(QFileDialog, "getOpenFileNames", return_value=(["/x/a(1-1).xlsx", "/x/a(1-2).xlsx"], "")):
            self.window.fs_response._pick()
        self.assertEqual(["/x/a(1-1).xlsx", "/x/a(1-2).xlsx"], self.window.fs_response.paths())
        self.assertIn("a(1-2).xlsx", self.window.fs_response.path_edit.toolTip())
        with patch.object(QFileDialog, "getOpenFileName", return_value=("/x/items.xlsx", "")):
            self.window.fs_iteminfo._pick()                            # single-file boxes are unchanged
        self.assertEqual(["/x/items.xlsx"], self.window.fs_iteminfo.paths())

    def test_folder_batch_fills_every_class_of_the_newest_exam_only(self):
        from PySide6.QtWidgets import QFileDialog, QMessageBox
        d = Path(self.tmp_dir())
        names = [f"{n}차 정기시험 교과목별 학생답 정오표({c}).xlsx" for n in (1, 2) for c in ("1-1", "1-2", "1-3")]
        names += [f"수행평가 강의실별 일람표({c}).xlsx" for c in ("1-1", "1-2")] + ["문항정보표(공통수학1) 샘플.xlsx"]
        for index, name in enumerate(names):
            path = d / name
            path.write_bytes(b"")
            os.utime(path, (1_000 + index, 1_000 + index))             # 2차 files are the newest
        with patch.object(QFileDialog, "getExistingDirectory", return_value=str(d)), \
                patch.object(QMessageBox, "information"), patch.object(QMessageBox, "warning"):
            self.window._pick_folder_batch()
        responses = [Path(p).name for p in self.window.fs_response.paths()]
        self.assertEqual([f"2차 정기시험 교과목별 학생답 정오표({c}).xlsx" for c in ("1-1", "1-2", "1-3")], responses)
        self.assertEqual(["수행평가 강의실별 일람표(1-1).xlsx", "수행평가 강의실별 일람표(1-2).xlsx"],
                         [Path(p).name for p in self.window.fs_perform.paths()])

    def test_folder_batch_reads_other_class_label_spellings(self):
        from PySide6.QtWidgets import QFileDialog, QMessageBox
        base = "1차 정기시험 교과목별 학생답 정오표"
        spellings = {
            "full-width parentheses": lambda c: f"（{c}）",
            "underscore": lambda c: f"_{c}",
            "en dash": lambda c: f"({c.replace('-', '–')})",
            "spaces inside": lambda c: f"( {c} )",
            "class word": lambda c: f"({c.split('-')[1]}반)",
        }
        for label, make in spellings.items():
            with self.subTest(label=label):
                d = Path(self.tmp_dir())
                for index, c in enumerate(("1-1", "1-2", "1-3")):
                    path = d / f"{base}{make(c)}.xlsx"
                    path.write_bytes(b"")
                    os.utime(path, (1_000 + index, 1_000 + index))
                with patch.object(QFileDialog, "getExistingDirectory", return_value=str(d)), \
                        patch.object(QMessageBox, "information") as info, patch.object(QMessageBox, "warning"):
                    self.window._pick_folder_batch()
                self.assertEqual(3, len(self.window.fs_response.paths()), label)
                self.assertNotIn("다른 반 정오표", info.call_args[0][2])

    def test_folder_batch_says_why_a_class_file_was_left_out(self):
        from PySide6.QtWidgets import QFileDialog, QMessageBox
        d = Path(self.tmp_dir())
        names = [f"1차 정기시험 교과목별 학생답 정오표({c}).xlsx" for c in ("1-1", "1-2")]
        names.append("1차 정기시험 교과목별 학생답 정오표(3학급).xlsx")       # label this app does not read
        for index, name in enumerate(names):
            path = d / name
            path.write_bytes(b"")
            os.utime(path, (1_000 + index, 1_000 + index))
        with patch.object(QFileDialog, "getExistingDirectory", return_value=str(d)), \
                patch.object(QMessageBox, "information") as info, patch.object(QMessageBox, "warning"):
            self.window._pick_folder_batch()
        self.assertEqual(2, len(self.window.fs_response.paths()))
        message = info.call_args[0][2]
        self.assertIn("같은 시험의 다른 반 정오표 1개", message)
        self.assertIn("1차 정기시험 교과목별 학생답 정오표(3학급).xlsx", message)

    def test_folder_batch_does_not_warn_about_another_round(self):
        from PySide6.QtWidgets import QFileDialog, QMessageBox
        d = Path(self.tmp_dir())
        names = [f"1차 정기시험 교과목별 학생답 정오표({c}).xlsx" for c in ("1-1", "1-2")]
        names += [f"2차 정기시험 교과목별 학생답 정오표({c}).xlsx" for c in ("1-1", "1-2")]
        for index, name in enumerate(names):
            path = d / name
            path.write_bytes(b"")
            os.utime(path, (1_000 + index, 1_000 + index))
        with patch.object(QFileDialog, "getExistingDirectory", return_value=str(d)), \
                patch.object(QMessageBox, "information") as info, patch.object(QMessageBox, "warning"):
            self.window._pick_folder_batch()
        self.assertEqual(2, len(self.window.fs_response.paths()))
        self.assertNotIn("다른 반 정오표", info.call_args[0][2])

    def tmp_dir(self):
        return self.resources.enter_context(tempfile.TemporaryDirectory())


if __name__ == "__main__":
    unittest.main()
