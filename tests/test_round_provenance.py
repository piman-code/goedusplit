from tests.runtime_isolation import ensure_isolated
ensure_isolated()
import unittest
from types import SimpleNamespace
from app.main_window import MainWindow
from app.portfolio_timeline import build_timeline

class RoundProvenanceTests(unittest.TestCase):
    def detect(self,response,manual='자동',header='2026학년도 1학기',perform='수행 2차.xlsx'):
        fake=SimpleNamespace(cmb_round=SimpleNamespace(currentText=lambda:manual),
            exam=SimpleNamespace(semester=header,source_files={'responses':response}),
            _round_file_names=lambda:{'정오표':response,'수행평가':perform,'분할점수':'','일람표':''})
        return MainWindow._detect_exam_round(fake)
    def test_response_first_round_beats_stale_second_round_selector(self):
        self.assertEqual((2026,1,1),self.detect('정오표 1차.xlsx','1학기 2차'))
    def test_performance_filename_cannot_create_missing_written_round(self):
        self.assertEqual((2026,1,None),self.detect('정오표.xlsx'))
    def test_unknown_legacy_stays_in_history_not_comparison_line(self):
        rows=[dict(student_hash='s',subject='수학',year=2026,semester_no=1,round_no=1,score=89.2),
              dict(student_hash='s',subject='수학',score=82)]
        model=build_timeline(rows,'s','수학')
        self.assertEqual([89.2],[r['score'] for r in model['points']])
        self.assertEqual(2,len(model['history']))
    def test_conflicting_written_inputs_stay_unknown(self):
        from app.data_loader import PATH_SEPARATOR
        names=PATH_SEPARATOR.join(['1학기 1차 정오표.xlsx','1학기 2차 정오표.xlsx'])
        self.assertEqual((2026,1,None),self.detect(names,'1학기 2차'))
    def test_partial_manual_only_fills_absent_component(self):
        self.assertEqual((2026,1,1),self.detect('1차 정오표.xlsx','2학기',header='2026학년도 1학기'))
        self.assertEqual((2026,2,1),self.detect('1차 정오표.xlsx','2학기',header='2026학년도'))
    def test_saved_work_uses_same_written_identity_policy(self):
        from app.portfolio_details import _term
        exam=SimpleNamespace(semester='2026학년도 1학기',source_files={'responses':'1차 정오표.xlsx'})
        state={'widgets':{'cmb_round':{'text':'1학기 2차'}},'inputs':{'fs_perform':'2차 수행.xlsx'}}
        self.assertEqual((2026,1,1),_term(exam,state))

class ActualAnalysisRoundTests(unittest.TestCase):
    def test_first_with_performance_then_distinct_second_retains_real_rounds(self):
        import tempfile
        from pathlib import Path
        from unittest.mock import patch
        import openpyxl
        from PySide6.QtWidgets import QApplication
        from app.synthetic_qa import analysis_fixtures
        from app.data_loader import PATH_SEPARATOR
        app=QApplication.instance() or QApplication([])
        root=Path(tempfile.mkdtemp(prefix='round-provenance-analysis-'))
        inputs=root/'inputs';inputs.mkdir()
        item,responses,performance=analysis_fixtures(inputs)
        second=openpyxl.load_workbook(responses[2][0])
        for row in (7,8):
            second.active.cell(row,4,'4');second.active.cell(row,5,'2')
            second.active.cell(row,6,0);second.active.cell(row,9,0)
        second.save(responses[2][0])
        with patch('app.main_window.QWebEngineView',None),patch.object(MainWindow,'_ai_material_root_dir',lambda self:root/'state'):
            w=MainWindow()
            try:
                w.fs_iteminfo.path_edit.setText(str(item))
                w.fs_perform.path_edit.setText(PATH_SEPARATOR.join(map(str,performance)))
                w.chk_perform.setChecked(True);w.spin_pencil_ratio.setValue(60);w.spin_perform_ratio.setValue(40)
                w.cmb_round.setCurrentText('1학기 2차')
                w.fs_response.path_edit.setText(PATH_SEPARATOR.join(map(str,responses[1])))
                w.run_analysis()
                self.assertEqual((2026,1,1),w._detect_exam_round())
                from app.portfolio_details import resolve_details
                first_points=w._load_portfolio_rows()
                combined=next(r for r in first_points if r['source_kind']=='current' and r['class_no']=='2/1')
                self.assertEqual(40.0,combined['score'])
                self.assertNotEqual(combined['score'],combined['pencil_total'])
                first=w._round_store().load_all()
                self.assertEqual(1,len(first));self.assertEqual(1,first[0]['round_no'])
                legacy=dict(first[0],round_no=2,label='1학기 2차')
                for key in ('round_provenance','response_fingerprint','append_revision'):
                    legacy.pop(key,None)
                w._round_store().save(legacy)
                w._refresh_rounds_tab()
                self.assertIsNone(w.cmb_round_other.currentData())
                self.assertEqual(0,w.table_round_metrics.rowCount())
                self.assertTrue(any('입력 출처 미확인' in w.cmb_round_other.itemText(i) for i in range(w.cmb_round_other.count())))
                preserved={p:p.read_bytes() for p in w._round_store().directory.glob('*.json')}
                w.fs_response.path_edit.setText(PATH_SEPARATOR.join(map(str,responses[2])))
                w.run_analysis()
                both=w._round_store().load_all()
                self.assertEqual([1,2],[r['round_no'] for r in both])
                self.assertNotEqual(both[0]['mean'],both[1]['mean'])
                self.assertNotEqual(both[0]['response_fingerprint'],both[1]['response_fingerprint'])
                self.assertEqual(5,len(w.exam.students));self.assertIsNotNone(w.perform_data)
                self.assertTrue(all(path.read_bytes()==data for path,data in preserved.items()))
                session=next(r for r in w._load_portfolio_rows() if r['source_kind']=='analysis-session' and r['round_no']==1 and r['class_no']=='2/1')
                details=resolve_details(w,session)
                self.assertTrue(details['tabs']['perform']['available'])
                self.assertEqual(40.0,session['score'])
                self.assertIn('전체 작업 저장 시 함께 보존',details['provenance'])
                before_cache=len(w._portfolio_analysis_sessions)
                w.fs_response.path_edit.clear()
                with patch('app.main_window.QMessageBox.warning'):
                    w.run_analysis()
                self.assertEqual(before_cache,len(w._portfolio_analysis_sessions))
            finally:w.close()
