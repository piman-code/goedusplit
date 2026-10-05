"""Opt-in native source-app analysis/export QA with freshly generated synthetic inputs.

Run: .venv/bin/python tests/manual_analysis_flows.py --run --output out_test/<new-folder>
No supplied student input, normal settings, installed app, or network is used.
This is source-app evidence; it does not certify Windows Excel or teacher acceptance.
"""
from pathlib import Path
import argparse
import csv
import hashlib
import json
import sys
from unittest.mock import patch

try:
    from tests.runtime_isolation import ensure_isolated
except ModuleNotFoundError:
    from runtime_isolation import ensure_isolated


def run(output):
    scratch = ensure_isolated(offscreen=False)
    from build_scripts.generate_synthetic_validation_kit import generate, _check_output, EXPECTED, IDENTITIES
    from PySide6.QtTest import QTest
    from PySide6.QtCore import QCoreApplication, QEvent
    from PySide6.QtWidgets import QApplication, QMessageBox, QPushButton
    import openpyxl
    import app.main_window as ui

    output = Path(output).absolute()
    _check_output(output)
    output.mkdir(parents=True, exist_ok=False)
    kit = generate(output / '합성 자료')
    checks = {}
    errors = []
    dialogs = {'real': False, 'option': QMessageBox.No, 'confirm': QMessageBox.No}
    destination = ['']
    folder = ['']
    input_path = ['']
    identifiers = [value for identity in IDENTITIES for value in identity]
    application = QApplication.instance() or QApplication([])

    def check(name, value):
        checks[name] = bool(value)
        if not value:
            raise AssertionError(name)

    def option(message):
        if message.checkBox():
            message.checkBox().setChecked(dialogs['real'])
        return dialogs['option']

    def warning(*args, **kwargs):
        if len(args) > 2 and args[2] == ui.REAL_IDENTITY_WARNING:
            return dialogs['confirm']
        errors.append('warning: ' + str(args[2]))
        return QMessageBox.No

    def hashes(root):
        return {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest()
                for p in root.rglob('*') if p.is_file()}

    original_inputs = hashes(kit)
    window = None
    report = {'source_app': True, 'frozen': False, 'synthetic_only': True,
              'checks': checks, 'errors': errors, 'status': 'incomplete'}
    with patch.object(ui.QFileDialog, 'getOpenFileName', side_effect=lambda *a, **k: (input_path[0], '')), \
         patch.object(ui.QFileDialog, 'getOpenFileNames', side_effect=lambda *a, **k: ([input_path[0]], '')), \
         patch.object(ui.QFileDialog, 'getSaveFileName', side_effect=lambda *a, **k: (destination[0], '')), \
         patch.object(ui.QFileDialog, 'getExistingDirectory', side_effect=lambda *a, **k: folder[0]), \
         patch.object(QMessageBox, 'exec', option), \
         patch.object(QMessageBox, 'warning', side_effect=warning), \
         patch.object(QMessageBox, 'critical', side_effect=lambda *a, **k: errors.append('critical: ' + str(a[2]))), \
         patch.object(QMessageBox, 'information', return_value=QMessageBox.Ok):
        try:
            window = ui.MainWindow()
            window.resize(1280, 800)
            window.show()
            QTest.qWait(600)
            for selector, filename in [(window.fs_iteminfo, '합성_문항정보표.xlsx'),
                                       (window.fs_response, '합성_정오표.xlsx')]:
                input_path[0] = str(kit / filename)
                selector.btn.click()
            window.btn_run_analysis.click()
            check('native file selection and written analysis',
                  window.exam is not None and len(window.exam.students) == 4
                  and [s.final_score for s in window.exam.students] == EXPECTED['exam']['written_scores']
                  and len(window.exam.items) == 3 and window.table_data.rowCount() == 4
                  and window.overall.mean == 46.875
                  and list(window.overall.levels_arr) == EXPECTED['exam']['written_levels']
                  and [s.p_value for s in window.item_stats] == [0.5, 0.5]
                  and window.overall.level_dist == {'A': 1, 'B': 0, 'C': 0, 'D': 0, 'E': 1, '미도달': 2})
            window.chk_perform.setChecked(True)
            input_path[0] = str(kit / '합성_수행평가.xlsx')
            window.fs_perform.btn.click()
            window.spin_pencil_ratio.setValue(60)
            window.spin_perform_ratio.setValue(40)
            window.btn_run_analysis.click()
            check('native performance analysis uses 60/40 independent totals',
                  [s.final_score for s in window.exam.students] == EXPECTED['exam']['combined_scores']
                  and window.overall.mean == 45.625)
            previous = (window.exam, window.overall)
            bad = output / '합성 invalid.xlsx'
            bad.write_text('synthetic invalid workbook only', encoding='utf-8')
            input_path[0] = str(bad)
            window.fs_response.btn.click()
            window.btn_run_analysis.click()
            expected_errors = len(errors)
            check('invalid input reports error and preserves existing analysis',
                  expected_errors == 1 and errors[0].startswith('critical:')
                  and window.exam is previous[0] and window.overall is previous[1])
            errors.clear()
            input_path[0] = str(kit / '합성_정오표.xlsx')
            window.fs_response.btn.click()

            csv_button = next(b for b in window.findChildren(QPushButton)
                              if b.text() == ui.SIDEBAR_CSV_EXPORT_BUTTON_TEXT)
            evidence_button = next(b for b in window.findChildren(QPushButton)
                                   if b.text() == ui.SIDEBAR_SPLITER_EXPORT_BUTTON_TEXT)
            csv_dir = output / '가명 CSV'
            csv_dir.mkdir()
            folder[0] = str(csv_dir)
            csv_button.click()
            with (csv_dir / '학생결과.csv').open(encoding='utf-8-sig', newline='') as f:
                rows = list(csv.DictReader(f))
            score_by_id = dict(zip(window._student_pseudonyms(), EXPECTED['exam']['combined_scores']))
            check('four real CSV exports preserve pseudonym-linked scores',
                  len(list(csv_dir.glob('*.csv'))) == 4 and len(rows) == 4
                  and len({row['가명 ID'] for row in rows}) == 4
                  and {row['가명 ID']: float(row['환산점수']) for row in rows} == score_by_id
                  and all(value not in p.read_text(encoding='utf-8-sig')
                          for value in identifiers for p in csv_dir.glob('*.csv')))
            evidence = output / '가명 근거.xlsx'
            destination[0] = str(evidence)
            dialogs['option'] = QMessageBox.Save
            evidence_button.click()
            book = openpyxl.load_workbook(evidence, data_only=False)
            student_rows = list(book['학생'].values)[1:]
            check('real evidence workbook sheets and pseudonym-linked scores',
                  {'요약', '성취수준 요약', '문항 근거', '학생', '학생 응답'}.issubset(book.sheetnames)
                  and len(student_rows) == 4
                  and len({row[0] for row in student_rows}) == 4
                  and {row[0]: row[5] for row in student_rows} == score_by_id
                  and all(value not in str(cell.value) for value in identifiers
                          for sheet in book for row in sheet for cell in row))
            book.close()

            before = hashes(output)
            dialogs['option'] = QMessageBox.Cancel
            csv_button.click()
            evidence_button.click()
            dialogs['option'] = QMessageBox.Save
            destination[0] = ''
            evidence_button.click()
            check('CSV/evidence option and save-dialog cancellation preserve all files', hashes(output) == before)
            dialogs.update(real=True, confirm=QMessageBox.No, option=QMessageBox.No)
            csv_button.click()
            dialogs['option'] = QMessageBox.Save
            destination[0] = str(output / 'declined.xlsx')
            evidence_button.click()
            check('declining real identity exports preserves all files', hashes(output) == before)

            real_dir = output / '승인한 합성 실명 CSV'
            real_dir.mkdir()
            folder[0] = str(real_dir)
            dialogs.update(option=QMessageBox.No, confirm=QMessageBox.Yes)
            original_name = window.exam.students[0].name
            window.exam.students[0].name = '=SYNTHETIC_FORMULA()'
            csv_button.click()
            with (real_dir / '학생결과.csv').open(encoding='utf-8-sig', newline='') as f:
                real_rows = list(csv.DictReader(f))
            check('explicit synthetic identity CSV escapes formula-like name',
                  len(real_rows) == 4 and real_rows[0]['이름'] == "'=SYNTHETIC_FORMULA()")
            dialogs['option'] = QMessageBox.Save
            destination[0] = str(output / '승인한 합성 실명.xlsx')
            evidence_button.click()
            book = openpyxl.load_workbook(destination[0], data_only=False)
            cell = book['학생'].cell(2, 4)
            check('explicit synthetic identity XLSX stores formula-like name as text',
                  cell.value == '=SYNTHETIC_FORMULA()' and cell.data_type == 's')
            book.close()
            window.exam.students[0].name = original_name
            window.save_current_subject_snapshot()
            snapshots = list(window._portfolio_store_dir().glob('*.json'))
            snapshot = json.loads(snapshots[0].read_text(encoding='utf-8')) if len(snapshots) == 1 else {}
            check('real snapshot saves scores and hashed identities only in isolated appdata',
                  window._portfolio_store_dir().is_relative_to(scratch)
                  and len(snapshot.get('students', [])) == 4
                  and sorted(s['final_score'] for s in snapshot['students']) == [0, 32.5, 50, 100]
                  and len({s['student_hash'] for s in snapshot['students']}) == 4
                  and all(value not in json.dumps(snapshot, ensure_ascii=False) for value in identifiers)
                  and all(set(s).isdisjoint({'sid', 'name', 'class_no'}) and len(s['student_hash']) == 24
                          for s in snapshot['students']))
            (output / 'snapshot-synthetic.json').write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding='utf-8')
            QTest.qWait(1200)  # Let deferred matplotlib/Qt paints complete before visual evidence.
            check('native analyzed window capture exists', window.grab().save(str(output / 'analysis-window.png')))
            check('all generated kit inputs and manifest remain unchanged', hashes(kit) == original_inputs)
            check('no unexpected native errors', not errors)
            report['status'] = 'passed'
        except Exception as error:
            report['status'] = 'failed'
            errors.append(type(error).__name__ + ': ' + str(error))
            if window:
                window.grab().save(str(output / 'failure-window.png'))
        finally:
            if window:
                window.close()
                QTest.qWait(700)
                window.deleteLater()
                QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
                application.processEvents()
            (output / 'ANALYSIS_QA.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2), flush=True)
    return 0 if report['status'] == 'passed' else 1


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run', action='store_true')
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    if not args.run:
        parser.error('--run is required; a native window appears with synthetic data')
    return run(args.output)


if __name__ == '__main__':
    raise SystemExit(main())
