#!/usr/bin/env python3
"""Create a new synthetic-only, locally verified manual validation kit."""
from __future__ import annotations

import argparse
from contextlib import contextmanager
from copy import deepcopy
import hashlib
import json
import math
import os
from pathlib import Path
import sys
import tempfile
from xml.sax.saxutils import escape
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
LEVELS = tuple("ABCDE")

# Independently authored fixtures. No source student document is read or copied.
IDENTITIES = [(f"합성-{n:03}", f"합성반/{n:02}", f"합성학생{n:02}") for n in range(1, 5)]
PAPER = """합성 검증 전용 시험지 - 실제 평가에 사용하지 않음
2026학년도 1학기 합성수학
선택형 2문항, 서답형 1문항
1. 합성 문항: 1+1은? [37.5점]
① 2 ② 3 ③ 4 ④ 5 ⑤ 6
2. 합성 문항: 2+2는? [37.5점]
① 2 ② 3 ③ 4 ④ 5 ⑤ 6
[서답형 1] 합성 문항: 풀이를 쓰시오. [25점]
"""
EXPECTED = {
    "synthetic_only": True,
    "derivation": {
        "written": "선택 37.5+37.5, 서답 25 = 100; 학생별 수기합 100/50/37.5/0",
        "performance": "만점 15+25=40; 합계 40/20/10/0 ÷40×100 = 100/50/25/0",
        "combined": "지필×0.6+수행환산×0.4 = 100/50/32.5/0",
        "calculator": "검토안별 직접 % 산술평균 → 배점가중합; NEIS만 묶음별 5% 반올림",
    },
    "exam": {
        "item_summary": [["선택형", 1, 37.5], ["선택형", 2, 37.5], ["서답형", 1, 25.0]],
        "student_ids": [f"합성-{n:03}" for n in range(1, 5)],
        "written_scores": [100.0, 50.0, 37.5, 0.0],
        "written_mean": 46.875,
        "choice_correct_matrix": [[1, 1], [1, 0], [0, 1], [0, 0]],
        "choice_rates": [0.5, 0.5],
        "written_levels": ["A", "E", "미도달", "미도달"],
        "performance_max": 40.0,
        "performance_scores": [100.0, 50.0, 25.0, 0.0],
        "combined_scores": [100.0, 50.0, 32.5, 0.0],
        "combined_mean": 45.625,
    },
    "paper": {
        "item_summary": [["선택형", 1, 37.5], ["선택형", 2, 37.5], ["서답형", 1, 25.0]],
        "choices": [5, 5, 0], "total_points": 100.0, "warnings": [],
    },
    "calculator": {
        "item_count": 2,
        "activeJudgeId": "synthetic-judge-2", "points": [1.5, 3.25],
        "mean_rates": [{"A": 100.0, "B": 80.0, "C": 63.25, "D": 35.0, "E": 0.0},
                       {"A": 100.0, "B": 75.0, "C": 55.0, "D": 35.0, "E": 15.0}],
        "total_points": 4.75,
        "raw_cuts": {"A": 4.75, "B": 3.6375, "C": 2.73625, "D": 1.6625, "E": 0.4875},
        "scaled_cuts": {"A": 100.0, "B": 76.57894736842105, "C": 57.60526315789474,
                        "D": 35.0, "E": 10.263157894736842},
        "neis_rates": {"A": 100, "B": 75, "C": 60, "D": 35, "E": 10},
        "neis_raw_cuts": {"A": 4.75, "B": 3.5625, "C": 2.85, "D": 1.6625, "E": 0.475},
        "neis_scaled_cuts": {"A": 100, "B": 75, "C": 60, "D": 35, "E": 10},
    },
    "coverage_limits": ["HWP binary not generated or validated", "scanned PDF not validated",
                        "minimal HWPX reader fixture, not a native Hancom compatibility certificate",
                        "synthetic data only; no school-PC or teacher acceptance evidence"],
}


def calculator_project():
    judge_rates = [((100, 82.5, 63.25, 40, 0), (100, 77.5, 63.25, 30, 0)),
                   ((100, 80, 60, 40, 20), (100, 70, 50, 30, 10))]
    project = {"version": 1, "judges": [{"id": f"synthetic-judge-{n}", "name": f"합성 검토안 {n}"}
                                         for n in (1, 2)],
               "activeJudgeId": "synthetic-judge-2", "evidenceMode": "difficultyAverage",
               "evidenceData": None, "items": []}
    for n, (points, rates) in enumerate(zip((1.5, 3.25), judge_rates), 1):
        judgments = {}
        for judge, values in zip(project["judges"], rates):
            judgments[judge["id"]] = {
                level: {"correct": [i < round(rate / 100 * 3) for i in range(3)],
                        "targetRate": rate, "overrideRate": rate}
                for level, rate in zip(LEVELS, values)}
        project["items"].append({"id": f"synthetic-item-{n}", "number": n,
                                 "title": f"합성 선택형 {n}번", "type": "선택형",
                                 "difficulty": "보통", "targetLevel": "C", "points": points,
                                 "sampleSize": 3, "standard": "합성 성취기준",
                                 "note": "합성 검증 전용", "evidence": [],
                                 "judgmentsByJudge": judgments})
    return project


def _linked(path):
    return path.is_symlink() or getattr(path, "is_junction", lambda: False)()


def _check_output(output):
    for path in [output, *output.parents]:
        if _linked(path):
            if (sys.platform == "darwin" and path in {Path("/var"), Path("/tmp"), Path("/etc")}
                    and path.is_symlink() and path.resolve() == Path("/private") / path.name):
                continue
            raise ValueError("Output or its parent is a symlink/junction; generation refused")
    if output.exists():
        raise FileExistsError("Output already exists; existing files are preserved")


@contextmanager
def _local_render_cache(output):
    with tempfile.TemporaryDirectory(prefix=".render-cache-", dir=output) as scratch:
        old = {key: os.environ.get(key) for key in ("MPLCONFIGDIR", "XDG_CACHE_HOME")}
        os.environ.update(MPLCONFIGDIR=scratch, XDG_CACHE_HOME=scratch)
        try:
            yield
        finally:
            for key, value in old.items():
                if value is None:
                    os.environ.pop(key, None)
                else:
                    os.environ[key] = value


def _workbooks(output):
    import openpyxl
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "합성 문항정보표"
    rows = [
        ["합성 검증 전용 (합성수학) 과목"], ["선택형 문항"],
        ["문항번호", "내용영역", "성취기준", "난이도", None, None, "배점", "정답"],
        [None, None, None, "어려움", "보통", "쉬움"],
        [1, "합성 연산", "[합성-01] 합성 기준", None, None, "○", 37.5, 1],
        [2, "합성 연산", "[합성-02] 합성 기준", None, "○", None, 37.5, 3],
        ["서답형 문항"],
        [1, "합성 풀이", "[합성-03] 합성 기준", "○", None, None, 25, "합성 정답"],
    ]
    for row in rows:
        sheet.append(row)
    book.save(output / "합성_문항정보표.xlsx")
    book.close()
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "합성 정오표"
    for row in [["합성 검증 전용 - 실제 학생자료 아님"], [],
                ["2026학년도 1학기 1학년 수학:합성수학"],
                ["반/번호", None, "성명", 1, 2, "선택형점수", "서답형점수", "기타점수", "과목총점"],
                [None, None, "정답", 1, 3], [None, None, "배점", 37.5, 37.5]]:
        sheet.append(row)
    for identity, values in zip(IDENTITIES, [(".", ".", 75, 25, 0, 100),
                                            (".", "2", 37.5, 12.5, 0, 50),
                                            ("4", ".", 37.5, 0, 0, 37.5), ("4", "2", 0, 0, 0, 0)]):
        sid, class_no, name = identity
        sheet.append([sid, class_no, name, *values])
    book.save(output / "합성_정오표.xlsx")
    book.close()
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "합성 수행평가"
    sheet.append(["합성 검증 전용"])
    sheet.append(["교과목 : 합성수학"])
    sheet.append(["반/번호", "합성 학번", "성명", "합성 영역가(만점 15.00,15.00%)",
                  "합성 영역나(만점 25.00,25.00%)", "합 계"])
    for identity, values in zip(IDENTITIES, [(15, 25, 40), (10, 10, 20), (5, 5, 10), (0, 0, 0)]):
        sid, class_no, name = identity
        sheet.append([class_no, sid, name, *values])
    book.save(output / "합성_수행평가.xlsx")
    book.close()


def _papers(output):
    (output / "합성_시험지.txt").write_text(PAPER, encoding="utf-8")
    paragraphs = "".join(f"<hp:p><hp:run><hp:t>{escape(line)}</hp:t></hp:run></hp:p>" for line in PAPER.splitlines())
    section = ('<?xml version="1.0" encoding="UTF-8"?>'
               '<hs:sec xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph" '
               'xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section">' + paragraphs + '</hs:sec>')
    with zipfile.ZipFile(output / "합성_시험지.hwpx", "x") as archive:
        archive.writestr("mimetype", "application/hwp+zip")
        archive.writestr("Contents/section0.xml", section)
    font = ROOT / "assets/fonts/NanumGothic.ttf"
    if any(_linked(path) for path in [font, *font.parents]) or not font.is_file():
        raise ValueError("Required bundled NanumGothic font missing or linked; no download attempted")
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_pdf import FigureCanvasPdf
    from matplotlib.font_manager import FontProperties
    import matplotlib
    with matplotlib.rc_context({"pdf.fonttype": 42}):
        figure = Figure(figsize=(8.27, 11.69))
        figure.text(0.07, 0.93, PAPER, va="top", fontsize=12,
                    fontproperties=FontProperties(fname=str(font)))
        FigureCanvasPdf(figure).print_pdf(str(output / "합성_시험지.pdf"),
                                         metadata={"Title": "합성 검증 시험지", "Creator": "Goedu-Split synthetic kit"})
        figure.clear()


def _same(actual, expected, label):
    if isinstance(expected, dict):
        for key, value in expected.items():
            _same(actual[key], value, label + "." + key)
    elif isinstance(expected, list):
        if len(actual) != len(expected):
            raise ValueError(label + ": length mismatch")
        for index, (left, right) in enumerate(zip(actual, expected)):
            _same(left, right, f"{label}[{index}]")
    elif isinstance(expected, (int, float)) and not isinstance(expected, bool):
        if not math.isclose(float(actual), expected, rel_tol=1e-12, abs_tol=1e-10):
            raise ValueError(label + ": arithmetic mismatch")
    elif actual != expected:
        raise ValueError(label + ": value mismatch")


def verify_created(output):
    """Read only the fixtures just generated, then compare with independent constants."""
    from app.analysis import analyze_items, analyze_overall, build_score_matrix, grade_level
    from app.data_loader import apply_perform, load_exam
    from app.perform_loader import load_perform
    from app.exam_structure import extract_exam_text, parse_exam_structure
    from app.expected_rates import build_neis_rows, summarize_designs
    checks = []
    exam = load_exam(output / "합성_문항정보표.xlsx", output / "합성_정오표.xlsx")
    summary = [[item.item_type, item.number, item.score] for item in exam.items]
    _same(summary, EXPECTED["exam"]["item_summary"], "exam.items")
    _same([student.sid for student in exam.students], EXPECTED["exam"]["student_ids"], "exam.identities")
    _same([student.total for student in exam.students], EXPECTED["exam"]["written_scores"], "exam.written")
    _same([grade_level(student.total, exam.cut_scores) for student in exam.students], EXPECTED["exam"]["written_levels"], "exam.levels")
    matrix, _, _ = build_score_matrix(exam)
    _same(matrix.tolist(), EXPECTED["exam"]["choice_correct_matrix"], "exam.matrix")
    _same([row.p_value for row in analyze_items(exam)[0]], EXPECTED["exam"]["choice_rates"], "exam.rates")
    _same(analyze_overall(exam).mean, EXPECTED["exam"]["written_mean"], "exam.mean")
    checks.append("NEIS-shaped synthetic XLSX and written analysis match handwritten constants")
    perform = load_perform(output / "합성_수행평가.xlsx")
    _same(perform.max_total, EXPECTED["exam"]["performance_max"], "perform.max")
    _same([perform.records[sid].pct100 for sid in EXPECTED["exam"]["student_ids"]], EXPECTED["exam"]["performance_scores"], "perform.scores")
    apply_perform(exam, perform, 60, 40)
    _same([student.final_score for student in exam.students], EXPECTED["exam"]["combined_scores"], "perform.combined")
    _same(analyze_overall(exam).mean, EXPECTED["exam"]["combined_mean"], "perform.mean")
    checks.append("Optional synthetic performance with 60/40 weights matches handwritten constants")
    for suffix in ("txt", "hwpx", "pdf"):
        text, method = extract_exam_text(output / f"합성_시험지.{suffix}")
        parsed = parse_exam_structure(text)
        actual = {"item_summary": [[item["type"], item["number"], item["points"]] for item in parsed["items"]],
                  "choices": [item["choices"] for item in parsed["items"]],
                  "total_points": parsed["total_points"], "warnings": parsed["warnings"]}
        _same(actual, EXPECTED["paper"], "paper." + suffix)
        checks.append(f"Synthetic Korean {suffix.upper()} parsed through {method}")
    project = json.loads((output / "합성_계산기_검토안.json").read_text(encoding="utf-8"))
    designs = []
    for item in project["items"]:
        rates = {level: sum(item["judgmentsByJudge"][judge["id"]][level]["overrideRate"]
                            for judge in project["judges"]) / 2 for level in LEVELS}
        designs.append({"number": item["number"], "type": item["type"], "difficulty": item["difficulty"],
                        "target": item["targetLevel"], "points": item["points"], "rates": rates})
    _same([design["rates"] for design in designs], EXPECTED["calculator"]["mean_rates"], "calculator.judge_average")
    _same(summarize_designs(designs), {key: EXPECTED["calculator"][key] for key in
          ("item_count", "total_points", "raw_cuts", "scaled_cuts", "neis_raw_cuts", "neis_scaled_cuts")}, "calculator.summary")
    rows = build_neis_rows(designs)
    _same({level: rows[0][level] for level in LEVELS}, EXPECTED["calculator"]["neis_rates"], "calculator.NEIS")
    checks.append("Synthetic review averages and decimal/0/100 rates match handwritten calculator constants")
    return {"status": "passed", "synthetic_only": True, "checks": checks,
            "limits": list(EXPECTED["coverage_limits"])}


def generate(output):
    output = Path(output).absolute()
    _check_output(output)
    output.mkdir(parents=True, exist_ok=False)
    with _local_render_cache(output):
        _workbooks(output)
        _papers(output)
        (output / "합성_읽어주세요.txt").write_text(
            "모든 파일은 독립 생성한 합성 검증자료입니다. 실제 학생자료가 아닙니다.\n"
            "사용 순서와 기대값: docs/SYNTHETIC_VALIDATION.md 및 expected.json\n"
            "HWP 바이너리·스캔 PDF·한글앱의 HWPX 편집 호환성은 검증하지 않았습니다.\n", encoding="utf-8")
        (output / "합성_계산기_검토안.json").write_text(json.dumps(calculator_project(), ensure_ascii=False, indent=2), encoding="utf-8")
        (output / "expected.json").write_text(json.dumps(deepcopy(EXPECTED), ensure_ascii=False, indent=2), encoding="utf-8")
        try:
            report = verify_created(output)
        except Exception as error:
            (output / "validation.json").write_text(json.dumps({"status": "failed", "synthetic_only": True,
                                                               "error": str(error)}, ensure_ascii=False, indent=2), encoding="utf-8")
            raise
        (output / "validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    from app import __version__
    files = {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in sorted(output.iterdir()) if path.is_file()}
    manifest = {"synthetic_only": True, "generator": "build_scripts/generate_synthetic_validation_kit.py",
                "app_version": __version__, "file_sha256": files,
                "hash_scope": "all generated files except this manifest; generation timestamps may differ between kits"}
    (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return output


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True, help="new output directory; use out_test for repository-local generation")
    args = parser.parse_args(argv)
    try:
        output = generate(args.output)
    except Exception as error:
        print("Synthetic kit generation failed: " + str(error), file=sys.stderr)
        return 1
    print("Synthetic kit verified: " + str(output))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
