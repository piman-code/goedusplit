"""Calculation contract for teacher estimates and the NEIS preparation table."""

from __future__ import annotations

import math
from pathlib import Path

LEVELS = ("A", "B", "C", "D", "E")
BOUNDARIES = ("A/B", "B/C", "C/D", "D/E", "E/미도달")
HEADERS = ["문항구분", "난이도", "해당문항번호", "문항수", "배점합", *LEVELS, "목표수준"]


def _number(value, label: str) -> float:
    if isinstance(value, bool) or value is None or value == "":
        raise ValueError(f"{label}: 숫자를 입력해 주세요.")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label}: 숫자를 입력해 주세요.") from exc
    if not math.isfinite(number):
        raise ValueError(f"{label}: 유한한 숫자를 입력해 주세요.")
    return number


def normalize_rates(rates) -> dict[str, float]:
    if not isinstance(rates, dict):
        raise ValueError("A~E 예상정답률을 모두 입력해 주세요.")
    result = {level: _number(rates.get(level), f"{level} 예상정답률") for level in LEVELS}
    if any(not 0 <= value <= 100 for value in result.values()):
        raise ValueError("예상정답률은 0~100%로 입력해 주세요.")
    return result


def validate_design_items(designs: list[dict]) -> None:
    if not isinstance(designs, list) or len(designs) > 1000:
        raise ValueError("문항 목록은 1,000개 이내여야 합니다.")
    seen = set()
    for index, item in enumerate(designs, 1):
        if not isinstance(item, dict):
            raise ValueError(f"{index}행 문항 형식을 확인해 주세요.")
        number = _number(item.get("number"), f"{index}행 문항번호")
        if number <= 0 or not number.is_integer():
            raise ValueError(f"{index}행 문항번호는 양의 정수여야 합니다.")
        item_type = item.get("type")
        if item_type not in {"선택형", "서답형"}:
            raise ValueError(f"{index}행 문항구분을 확인해 주세요.")
        key = (item_type, number)
        if key in seen:
            raise ValueError(f"{item_type} {int(number)}번이 중복되었습니다.")
        seen.add(key)
        points = _number(item.get("points"), f"{index}행 배점")
        if not 0 < points <= 10000:
            raise ValueError(f"{index}행 배점은 0보다 크고 10,000점 이하여야 합니다.")
        if item.get("difficulty") not in {"쉬움", "보통", "어려움"}:
            raise ValueError(f"{index}행 난이도를 확인해 주세요.")
        rates = normalize_rates(item.get("rates"))
        if any(rates[upper] < rates[lower] for upper, lower in zip(LEVELS, LEVELS[1:])):
            raise ValueError(f"{item_type} {int(number)}번: A~E 예상정답률의 역전을 확인해 주세요.")


def round_neis_rate(value: float) -> int:
    number = _number(value, "NEIS 예상정답률")
    if not 0 <= number <= 100:
        raise ValueError("NEIS 예상정답률은 0~100%여야 합니다.")
    # Match JavaScript's nonnegative half-up rounding at a binary-float tie.
    return min(100, 5 * math.floor(number / 5 + 0.5 + 1e-12))


def build_neis_rows(designs: list[dict]) -> list[dict]:
    validate_design_items(designs)
    groups = {}
    for item in designs:
        groups.setdefault((item["type"], item["difficulty"]), []).append(item)
    type_order = {"선택형": 0, "서답형": 1}
    diff_order = {"쉬움": 0, "보통": 1, "어려움": 2}
    rows = []
    for key in sorted(groups, key=lambda k: (type_order[k[0]], diff_order[k[1]])):
        items = sorted(groups[key], key=lambda i: int(i["number"]))
        total = math.fsum(float(i["points"]) for i in items)
        row = {
            "문항구분": key[0], "난이도": key[1],
            "해당문항번호": ", ".join(str(int(i["number"])) for i in items),
            "문항수": len(items), "배점합": total,
            "목표수준": ", ".join(f"{int(i['number'])}:{i.get('target', '')}" for i in items),
        }
        for level in LEVELS:
            weighted = math.fsum(float(i["points"]) * float(i["rates"][level]) for i in items) / total
            row[level] = round_neis_rate(weighted)
        rows.append(row)
    return rows


def summarize_designs(designs: list[dict]) -> dict:
    rows = build_neis_rows(designs)
    total = math.fsum(float(i["points"]) for i in designs)
    raw = {level: math.fsum(float(i["points"]) * float(i["rates"][level]) / 100 for i in designs) for level in LEVELS}
    neis = {level: math.fsum(row["배점합"] * row[level] / 100 for row in rows) for level in LEVELS}
    return {
        "item_count": len(designs), "total_points": total,
        "raw_cuts": raw, "scaled_cuts": {lv: raw[lv] / total * 100 if total else 0 for lv in LEVELS},
        "neis_raw_cuts": neis, "neis_scaled_cuts": {lv: neis[lv] / total * 100 if total else 0 for lv in LEVELS},
    }


def write_estimation_workbook(path: str | Path, designs: list[dict]) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    summary = summarize_designs(designs)
    if not designs:
        raise ValueError("저장할 문항이 없습니다.")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "분할점수 비교"
    sheet.append(["문항수", summary["item_count"], "총점", summary["total_points"]])
    sheet.append(["경계", "문항별 원점수", "문항별 100점 환산", "NEIS 표 원점수", "NEIS 표 100점 환산", "환산 차이"])
    for level, boundary in zip(LEVELS, BOUNDARIES):
        sheet.append([boundary, summary["raw_cuts"][level], summary["scaled_cuts"][level], summary["neis_raw_cuts"][level], summary["neis_scaled_cuts"][level], summary["neis_scaled_cuts"][level] - summary["scaled_cuts"][level]])
    sheet.append([])
    sheet.append(["문항별 원점수 = 전체 문항의 배점 × 예상정답률 합계. 100점 환산 = 원점수 ÷ 총점 × 100."])
    sheet.append(["NEIS 표는 문항구분·난이도별 배점가중평균을 구한 뒤 A~E 모두 5% 단위로 반올림한 준비표입니다."])
    sheet.append(["목표수준 기본값은 조정용 출발점입니다. 학교의 성취수준 준거와 교과협의를 거쳐 확정하세요."])
    details = workbook.create_sheet("문항별 입력값")
    details.append(["문항구분", "문항", "배점", "난이도", "목표수준", *LEVELS, "성취기준"])
    for item in designs:
        details.append([item["type"], int(item["number"]), float(item["points"]), item["difficulty"], item.get("target", ""), *[float(item["rates"][lv]) for lv in LEVELS], str(item.get("standard", ""))])
    neis_sheet = workbook.create_sheet("NEIS 입력표")
    neis_sheet.append(HEADERS)
    for row in build_neis_rows(designs):
        neis_sheet.append([row.get(header, "") for header in HEADERS])
    for page in workbook:
        page.freeze_panes = "A3" if page is sheet else "A2"
        for cells in page:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                else:
                    cell.number_format = "0.00########"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        header_row = 2 if page is sheet else 1
        for cell in page[header_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEDEA")
        for column in page.columns:
            page.column_dimensions[column[0].column_letter].width = 22
    workbook.save(path)
