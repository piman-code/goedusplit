"""Compare teacher estimates with how students actually did (예측-실측 보정 리포트).

Teachers estimate the correct rate of the *minimally competent* student of each
level, i.e. a student right at that level's cut score. The observable group is
the students nearest the cut score, the same number from each side. A group taken only from
inside the level (e.g. its lowest third) sits above the cut and would make
accurate estimates look too low. The level-wide mean is kept as a reference.

The overall level of the estimates cannot be checked this way: students at a
cut score score about that cut on average, so summed over all items their
rates reproduce the applied cut. The level-wide difference therefore shows
"estimated cut vs applied cut". What the data does show is which items and
difficulties were over- or under-estimated relative to the rest (``relative``).
"""

from __future__ import annotations

import math
from pathlib import Path

LEVELS = ("A", "B", "C", "D", "E")
BOUNDARIES = ("A/B", "B/C", "C/D", "D/E", "E/미도달")
BORDER_SHARE = 0.10       # of all students, per cut score
BORDER_MIN = 3
BORDER_MAX_DISTANCE = 10.0  # points on the 100-point final score
LARGE_GAP = 15.0            # %p
FEW_STUDENTS = 5
DIFFICULTY_DELTA = {"쉬움": 10, "보통": 0, "어려움": -10}  # what the app adds to a preset row per difficulty


def borderline_indices(scores: list[float], cuts: dict) -> dict[str, list[list[int]]]:
    """Students nearest each level's cut score, as [just below, at or above].

    About 10% of all students (at least 3) per cut, half from each side,
    keeping ties at the last distance on each side and ignoring students more
    than 10 points away. Rates average the two sides equally, so the estimate
    stays centred on the cut even where one side is crowded, as near the top or
    bottom of the distribution.
    """
    per_side = math.ceil(max(BORDER_MIN, math.ceil(len(scores) * BORDER_SHARE)) / 2)
    result = {}
    for level in LEVELS:
        cut = cuts.get(level)
        if cut is None:
            continue
        sides = []
        for side in (lambda s: s < cut, lambda s: s >= cut):
            near = sorted(
                (i for i, s in enumerate(scores) if side(float(s)) and abs(float(s) - float(cut)) <= BORDER_MAX_DISTANCE),
                key=lambda i: (abs(float(scores[i]) - float(cut)), i),
            )
            if near:
                limit = abs(float(scores[near[min(per_side, len(near)) - 1]]) - float(cut))
                sides.append([i for i in near if abs(float(scores[i]) - float(cut)) <= limit])
        if sides:
            result[level] = sides
    return result


def _mean_pct(values: list[float]) -> float | None:
    return math.fsum(values) / len(values) * 100 if values else None


def _border_pct(per_student: list[float], sides: list[list[int]]) -> float | None:
    means = [_mean_pct([per_student[i] for i in side]) for side in sides]
    return math.fsum(means) / len(means) if means else None


def _level_rates(per_student: list[float], levels: list[str], border: dict[str, list[list[int]]]) -> dict:
    """per_student: 0~1 score rate of one item (or item group) for every student."""
    return {
        "all": {lv: _mean_pct([per_student[i] for i, x in enumerate(levels) if x == lv]) for lv in LEVELS},
        "border": {lv: _border_pct(per_student, border.get(lv, [])) for lv in LEVELS},
    }


def _diffs(predicted: dict, actual: dict) -> dict:
    return {lv: None if actual[lv] is None else actual[lv] - predicted[lv] for lv in LEVELS}


def _weighted_mean(pairs) -> float | None:
    """pairs: (weight, value); None values are skipped."""
    pairs = [(w, v) for w, v in pairs if v is not None]
    total = math.fsum(w for w, _ in pairs)
    return math.fsum(w * v for w, v in pairs) / total if total else None


def observed_target(row: dict, rates_for, min_levels: int = 3) -> str | None:
    """The target level whose default rates are closest to how this item's borderline students did.

    ``rates_for(target, difficulty)`` is the app's preset row for that difficulty, the same
    rates a new item with that target starts from, so the chosen level reproduces the real
    A~E pattern best (least squares over the boundaries with students). A difficulty label
    alone does not: in real results 쉬움 items ranged from target A to E. None when fewer
    than ``min_levels`` boundaries have borderline students. A 서답형 group row carries
    ``members`` [(points, difficulty)] and is compared with their 배점-weighted default rates,
    since each item later starts from the row for its own difficulty.
    """
    have = [lv for lv in LEVELS if row["border"][lv] is not None]
    if len(have) < min_levels:
        return None
    members = row.get("members") or [(1.0, row["difficulty"])]
    weight = math.fsum(points for points, _ in members)
    rows = {}  # rates_for reads the app's saved table: once per target and difficulty

    def default_rate(target, level):
        total = 0.0
        for points, difficulty in members:
            key = (target, difficulty if difficulty in DIFFICULTY_DELTA else "보통")
            if key not in rows:
                rows[key] = rates_for(*key)
            total += points * float(rows[key][level])
        return total / weight

    def error(target):
        return math.fsum((row["border"][lv] - default_rate(target, lv)) ** 2 for lv in have)

    return min(LEVELS, key=error)


def observed_targets(report: dict) -> dict[tuple[str, int], str]:
    """{(type, number): observed target}. 서답형 items share their group's (only totals are known)."""
    result = {}
    for row in report["rows"]:
        target = row.get("observed_target")
        if target is None:
            continue
        kind = "서답형" if row["type"] == "서답형 묶음" else row["type"]
        for number in str(row["number"]).split(","):
            result[(kind, int(number))] = target
    return result


def build_calibration(designs: list[dict], exam, levels: list[str], rates_for=None) -> dict:
    """designs: calculator items with type, number, difficulty, target, points, rates and source_item.

    With ``rates_for`` every row also gets ``observed_target`` (see observed_target).
    """
    students = exam.students
    scores = [float(st.final_score) for st in students]
    border = borderline_indices(scores, exam.cut_scores)
    rows, unmatched, notes = [], [], []
    select = [d for d in designs if d["type"] == "선택형"]
    serdap = [d for d in designs if d["type"] == "서답형"]

    for design in select:
        item = design.get("source_item")
        if item is None:
            unmatched.append(f"선택형 {int(design['number'])}번")
            continue
        per_student = [1.0 if str(st.answers.get(item.number, "")).strip() == "." else 0.0 for st in students]
        rows.append(_row(design["type"], str(int(design["number"])), design["difficulty"], design.get("target", ""),
                         float(design["points"]), design["rates"], _level_rates(per_student, levels, border)))

    unmatched.extend(f"서답형 {int(d['number'])}번" for d in serdap if d.get("source_item") is None)
    serdap = [d for d in serdap if d.get("source_item") is not None]
    exam_serdap = sorted(int(it.number) for it in exam.items if it.item_type == "서답형")
    serdap_max = math.fsum(float(it.score) for it in exam.items if it.item_type == "서답형")
    if serdap and sorted(int(d["number"]) for d in serdap) != exam_serdap:
        notes.append("서답형은 정오표에 학생별 총점만 있어, 시험의 모든 서답형 문항에 예측값이 있어야 묶음으로 비교합니다. 이번에는 서답형을 비교하지 않았습니다.")
    elif serdap and serdap_max > 0:
        points = math.fsum(float(d["points"]) for d in serdap)
        predicted = {lv: math.fsum(float(d["points"]) * float(d["rates"][lv]) for d in serdap) / points for lv in LEVELS}
        per_student = [min(1.0, max(0.0, float(st.serdap_score) / serdap_max)) for st in students]
        numbers = ", ".join(str(int(d["number"])) for d in sorted(serdap, key=lambda d: int(d["number"])))
        rows.append(_row("서답형 묶음", numbers, "", "", points, predicted, _level_rates(per_student, levels, border)))
        rows[-1]["members"] = [(float(d["points"]), d["difficulty"]) for d in serdap]

    if getattr(exam, "use_perform", False) and float(getattr(exam, "weight_perform", 0) or 0) > 0:
        notes.append("성취수준은 수행평가를 합산한 환산점수로 나뉘었고, 정답률은 지필 결과입니다. 수행평가 비중만큼 차이가 생길 수 있습니다.")

    designed = {(d["type"], int(d["number"])) for d in designs}
    not_designed = [f"{it.item_type} {it.number}번" for it in exam.items if (it.item_type, int(it.number)) not in designed]

    for row in rows:
        row["observed_target"] = observed_target(row, rates_for) if rates_for else None

    offset = {lv: _weighted_mean((row["points"], row["diff"][lv]) for row in rows) for lv in LEVELS}
    for row in rows:
        row["relative"] = {
            lv: None if row["diff"][lv] is None or offset[lv] is None else row["diff"][lv] - offset[lv] for lv in LEVELS
        }

    total = math.fsum(row["points"] for row in rows)
    cuts = []
    for level, boundary in zip(LEVELS, BOUNDARIES):
        predicted = math.fsum(row["points"] * row["predicted"][level] / 100 for row in rows)
        usable = all(row["border"][level] is not None for row in rows) and bool(rows)
        actual = math.fsum(row["points"] * row["border"][level] / 100 for row in rows) if usable else None
        applied = exam.cut_scores.get(level)
        cuts.append({
            "boundary": boundary, "level": level,
            "predicted_scaled": predicted / total * 100 if total else None,
            "applied": None if applied is None else float(applied),
            "actual_scaled": actual / total * 100 if (usable and total) else None,
        })
    counts = {lv: sum(1 for x in levels if x == lv) for lv in LEVELS}
    return {
        "rows": rows,
        "unmatched": unmatched,
        "not_designed": not_designed,
        "counts": counts,
        "border_counts": {lv: sum(len(side) for side in border.get(lv, [])) for lv in LEVELS},
        "border_one_sided": [lv for lv in LEVELS if len(border.get(lv, [])) == 1],
        "notes": notes,
        "offset": offset,
        "relative_by_difficulty": {
            diff: {
                "items": sum(1 for row in rows if row["difficulty"] == diff),
                "levels": {lv: _weighted_mean((row["points"], row["relative"][lv]) for row in rows if row["difficulty"] == diff) for lv in LEVELS},
            }
            for diff in ("쉬움", "보통", "어려움") if any(row["difficulty"] == diff for row in rows)
        },
        "large_gaps": sum(1 for row in rows for lv in LEVELS if row["relative"][lv] is not None and abs(row["relative"][lv]) >= LARGE_GAP),
        "cuts": cuts,
        "total_points": total,
    }


def _row(kind, number, difficulty, target, points, predicted, actual) -> dict:
    predicted = {lv: float(predicted[lv]) for lv in LEVELS}
    return {
        "type": kind, "number": number, "difficulty": difficulty, "target": target, "points": points,
        "predicted": predicted, "border": actual["border"], "all": actual["all"],
        "diff": _diffs(predicted, actual["border"]),
    }


def _direction(value: float) -> str:
    return "높게" if value < 0 else "낮게"


def summary_lines(report: dict) -> list[str]:
    """What the teacher should read first: relative misses by difficulty, then cautions."""
    lines = []
    groups = report["relative_by_difficulty"]
    if len(groups) < 2:
        # Relative misses average to zero, so a single difficulty group always looks "on target".
        lines.append("문항 난이도가 한 가지뿐이라 난이도별 경향은 알 수 없습니다. 표에서 상대차가 큰 문항을 확인하세요.")
        groups = {}
    for diff, data in groups.items():
        values = [(lv, v) for lv, v in data["levels"].items() if v is not None]
        if not values:
            continue
        mean = math.fsum(v for _, v in values) / len(values)
        detail = ", ".join(f"{lv} {v:+.0f}" for lv, v in values)
        if abs(mean) < 2.5:
            lines.append(f"{diff} 문항({data['items']}개): 다른 문항과 비슷하게 예측했습니다 (경계별 실측-예측 상대차 {detail}).")
        else:
            lines.append(f"{diff} 문항({data['items']}개): 다른 문항보다 경계 학생 정답률을 평균 {abs(mean):.1f}%p {_direction(mean)} 예측했습니다 ({detail}).")
    differ = [row for row in report["rows"] if row["type"] == "선택형" and row.get("observed_target")
              and row["target"] in LEVELS and row["observed_target"] != row["target"]]
    if differ:
        shown = ", ".join(f"{row['number']}번 {row['target']}→{row['observed_target']}" for row in differ[:6])
        more = f" 외 {len(differ) - 6}문항" if len(differ) > 6 else ""
        lines.append(f"목표수준이 실측과 다른 선택형 {len(differ)}문항: {shown}{more} (실측 목표 = 경계 학생 결과에 가장 가까운 기준표 줄).")
    if report["large_gaps"]:
        lines.append(f"다른 문항보다 {LARGE_GAP:.0f}%p 이상 빗나간 칸이 {report['large_gaps']}개 있습니다(표에서 색 표시).")
    for level, boundary in zip(LEVELS, BOUNDARIES):
        if report["offset"][level] is None:
            lines.append(f"{boundary}: 분할점수 가까이에 학생이 없어 비교할 수 없습니다.")
            continue
        cautions = []
        if report["border_counts"][level] < FEW_STUDENTS:
            cautions.append("경계 학생이 적음")
        if level in report.get("border_one_sided", []):
            cautions.append("분할점수 한쪽에만 학생이 있음")
        if cautions:
            lines.append(f"{boundary}: {', '.join(cautions)} — 참고만 하세요.")
    return lines


def headline_lines(report: dict, count: int = 3) -> list[str]:
    """What to look at first, in plain words: the cells missed most relative to other items."""
    cells = [
        (abs(row["relative"][lv]), row, lv, boundary)
        for row in report["rows"] for lv, boundary in zip(LEVELS, BOUNDARIES)
        if row["relative"][lv] is not None and abs(row["relative"][lv]) >= LARGE_GAP
    ]
    if not cells:
        return ["크게 빗나간 칸이 없습니다. 이번 예측은 문항끼리 고르게 맞았습니다."]
    cells.sort(key=lambda cell: -cell[0])
    lines = []
    for _, row, lv, boundary in cells[:count]:
        kind = row["type"] if row["type"] != "서답형 묶음" else "서답형 묶음"
        lines.append(
            f"{kind} {row['number']}번 · {boundary} 경계: 예측 {row['predicted'][lv]:.0f}% → 실제 {row['border'][lv]:.0f}% "
            f"(다른 문항보다 {abs(row['relative'][lv]):.0f}%p {_direction(row['relative'][lv])} 예측)"
        )
    more = len(cells) - count
    tail = f" 그 밖에 {more}칸이 더 있습니다(표에서 색 칸)." if more > 0 else ""
    lines.append("→ 이 문항들의 목표수준과 예상정답률을 다시 보세요. 다음 시험 기본값은 '다음 시험 기준표 제안'으로 고칠 수 있습니다." + tail)
    return lines


def cut_lines(report: dict) -> list[str]:
    """The level-wide difference is the estimated cut against the applied cut, not estimate quality."""
    if report["not_designed"] or report["unmatched"]:
        return ["일부 문항만 예측값이 있어 검토안 분할점수를 이번 적용 분할점수와 직접 비교하지 않았습니다."]
    if any("수행평가" in note for note in report["notes"]):
        return ["수행평가를 합산한 분할점수라 지필 검토안 분할점수와 직접 비교하지 않았습니다."]
    lines = []
    for cut in report["cuts"]:
        if cut["predicted_scaled"] is None or cut["applied"] is None:
            continue
        gap = cut["predicted_scaled"] - cut["applied"]
        if abs(gap) >= 0.5:
            lines.append(f"{cut['boundary']}: 검토안으로 계산한 분할점수 {cut['predicted_scaled']:.1f}점, 이번에 적용한 분할점수 {cut['applied']:.1f}점.")
    return lines


def suggest_presets(report: dict, current: dict[str, dict], min_items: int = 2) -> dict[str, dict]:
    """Next exam's default rates per target level, from this exam's borderline students.

    For every target level with at least ``min_items`` selection items, the
    suggestion is the mean borderline rate of those items at each boundary.
    Preset rows are for 보통 items and the app adds DIFFICULTY_DELTA when it uses
    them, so each item's difficulty adjustment is taken off before averaging.
    Cells without data keep the current value.
    """
    result = {}
    for target in LEVELS:
        rows = [row for row in report["rows"] if row["type"] == "선택형" and row["target"] == target]
        suggested = None
        if len(rows) >= min_items:
            suggested = {}
            for level in LEVELS:
                values = [row["border"][level] - DIFFICULTY_DELTA.get(row["difficulty"], 0)
                          for row in rows if row["border"][level] is not None]
                value = math.floor(math.fsum(values) / len(values) + 0.5) if values else current[target][level]
                suggested[level] = max(0, min(100, value))
        result[target] = {"items": len(rows), "current": dict(current[target]), "suggested": suggested}
    return result


def write_calibration_workbook(path: str | Path, report: dict) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "요약"
    summary.append(["예측-실측 보정 리포트. 실측 = 각 분할점수 위아래로 가장 가까운 학생(경계 학생)의 정답률."])
    summary.append(["상대차 = 문항의 실측-예측에서 그 경계의 배점 가중 평균 차이를 뺀 값. 문항끼리의 예측 정확도를 봅니다."])
    summary.append(["학생 이름·학번은 들어가지 않지만, 경계 학생이 적은 수준의 값은 소수 학생의 결과이므로 공유에 주의하세요."])
    summary.append([])
    summary.append(["경계", "수준 학생 수", "경계 학생 수", "검토안 분할점수", "적용 분할점수", "경계 학생 실측 분할점수"])
    for level, cut in zip(LEVELS, report["cuts"]):
        summary.append([cut["boundary"], report["counts"][level], report["border_counts"][level],
                        cut["predicted_scaled"], cut["applied"], cut["actual_scaled"]])
    summary.append([])
    for line in summary_lines(report) + cut_lines(report) + report["notes"]:
        summary.append([line])
    if report["unmatched"]:
        summary.append(["분석 자료에서 찾지 못한 예측 문항: " + ", ".join(report["unmatched"])])
    if report["not_designed"]:
        summary.append(["예측값이 없는 시험 문항: " + ", ".join(report["not_designed"])])

    items = workbook.create_sheet("문항별")
    headers = ["구분", "번호", "난이도", "목표수준", "실측 목표수준", "배점"]
    for level in LEVELS:
        headers += [f"{level} 예측", f"{level} 실측(경계)", f"{level} 차이", f"{level} 상대차", f"{level} 수준 평균"]
    items.append(headers)
    for row in report["rows"]:
        values = [row["type"], row["number"], row["difficulty"], row["target"], row.get("observed_target") or "", row["points"]]
        for level in LEVELS:
            values += [row["predicted"][level], row["border"][level], row["diff"][level], row["relative"][level], row["all"][level]]
        items.append(values)
    for sheet in workbook:
        for cells in sheet.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.0"
    for cell in items[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEDEA")
    items.freeze_panes = "G2"
    workbook.save(path)
